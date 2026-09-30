"""Genera la presentación de la Tarea 5 (WTAP con atención y puntero) a partir de figuras_t5/.

Uso (desde la raíz del repositorio):
    python scripts/make_presentation_t5.py
    soffice --headless --convert-to pdf Tarea5_WTAP_presentacion.pptx
"""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figuras_t5")
OUT = os.path.join(ROOT, "Tarea5_WTAP_presentacion.pptx")
R = json.load(open(os.path.join(FIG, "resultados.json"), encoding="utf-8"))

FONT = "Arial"
NAVY = RGBColor(0x1B, 0x2A, 0x41)
INK = RGBColor(0x2E, 0x35, 0x40)
MUTED = RGBColor(0x6B, 0x75, 0x85)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
ORANGE = RGBColor(0xFF, 0x7F, 0x0E)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
PURPLE = RGBColor(0x94, 0x67, 0xBD)
RED = RGBColor(0xC0, 0x39, 0x2B)
PALE_BLUE = RGBColor(0xDC, 0xEA, 0xF7)
PALE_RED = RGBColor(0xF9, 0xDE, 0xDC)
PALE_GREEN = RGBColor(0xDF, 0xF0, 0xDF)
N_SLIDES = 5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers (los mismos de la Tarea 4)
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, shape=None):
    """runs: str, o lista de párrafos; cada párrafo es str o lista de (texto, {opciones})."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = runs if isinstance(runs, list) else [runs]
    for k, par in enumerate(paragraphs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        pieces = par if isinstance(par, list) else [(par, {})]
        for txt, opt in pieces:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = opt.get("font", FONT)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
        if isinstance(par, list) and par and "space_after" in par[0][1]:
            p.space_after = Pt(par[0][1]["space_after"])
    return shape


def image(slide, path, x, y, w=None, h=None, align="center"):
    """Inserta la imagen ajustada a la caja (x, y, w, h) manteniendo la proporción."""
    iw, ih = Image.open(path).size
    ratio = iw / ih
    if w is not None and h is not None:
        if w / h > ratio:
            nw, nh = h * ratio, h
        else:
            nw, nh = w, w / ratio
        ox = x + {"center": (w - nw) / 2, "left": 0, "right": w - nw}[align]
        oy = y + (h - nh) / 2
    elif w is not None:
        nw, nh, ox, oy = w, w / ratio, x, y
    else:
        nw, nh, ox, oy = h * ratio, h, x, y
    return slide.shapes.add_picture(path, Inches(ox), Inches(oy), Inches(nw), Inches(nh))


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def header(slide, n, title, subtitle=None):
    box(slide, 0, 0, SW, SH, fill=WHITE)
    box(slide, 0.5, 0.42, 0.09, 0.62, fill=BLUE)
    text(slide, 0.72, 0.3, 11.8, 0.6, title, size=27, bold=True, color=NAVY)
    if subtitle:
        text(slide, 0.74, 0.86, 11.9, 0.4, subtitle, size=14, color=MUTED)
    text(slide, 0.5, 7.05, 8, 0.3, "OII462 · Tarea 5 · Atención con puntero para el WTAP", size=10, color=MUTED)
    text(slide, SW - 1.3, 7.05, 0.8, 0.3, f"{n} / {N_SLIDES}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=14, body_size=12, fill=LIGHT):
    box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(slide, x, y + 0.15, 0.07, h - 0.3, fill=accent)
    text(slide, x + 0.22, y + 0.1, w - 0.35, 0.4, title, size=title_size, bold=True, color=accent)
    text(slide, x + 0.22, y + 0.48, w - 0.35, h - 0.55, body, size=body_size, color=INK, spacing=1.1)


def kpi(slide, x, y, w, h, value, label, color, sub=None):
    box(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(slide, x, y, w, 0.08, fill=color)
    text(slide, x, y + 0.12, w, 0.55, value, size=25, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(slide, x + 0.05, y + 0.66, w - 0.1, 0.35, label, size=11.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
    if sub:
        text(slide, x + 0.05, y + 0.94, w - 0.1, 0.3, sub, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def formula_png(latex, path, fontsize=26, color="#1B2A41"):
    fig = plt.figure(figsize=(8, 1))
    fig.text(0, 0.5, latex, fontsize=fontsize, color=color, va="center")
    fig.savefig(path, dpi=250, bbox_inches="tight", pad_inches=0.05, transparent=True)
    plt.close(fig)
    return path


def chip(slide, x, y, w, h, label, fill, size=12, color=WHITE, bold=True):
    s = box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    text(slide, 0, 0, 0, 0, label, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=s)
    return s


def pct(v, d=1):
    return f"{v:.{d}f}%".replace(".", ",")


def miles(n):
    return f"{n:,}".replace(",", ".")


def num(v, d=3):
    return f"{v:.{d}f}".replace(".", ",")


MONO = {"font": "Courier New"}
T_ = R["test"]
MOD, MMR, EXP = "Modelo (atención)", "MMR (heurística clásica)", "Experto (ILS)"
D = R["demo"]
tmp = os.path.join(FIG, "_tmp")
os.makedirs(tmp, exist_ok=True)

# ================================================================ 1. Título + decisión puntual
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=WHITE)
box(s, 0, 0, 4.7, SH, fill=NAVY)
text(s, 0.55, 0.7, 3.8, 0.4, "OII462 · TAREA 5", size=13, bold=True, color=RGBColor(0x8F, 0xB8, 0xDE))
text(s, 0.55, 1.15, 3.9, 2.4, "Atención con puntero para el WTAP", size=32, bold=True, color=WHITE, spacing=1.0)
text(s, 0.55, 2.95, 3.8, 1.6, "Un encoder con self-attention y un puntero que aprenden a asignar armas a objetivos "
     "imitando paso a paso a un experto (ILS)", size=15, color=RGBColor(0xD6, 0xE2, 0xEE), spacing=1.15)
box(s, 0.55, 4.95, 0.9, 0.05, fill=ORANGE)
text(s, 0.55, 5.15, 3.8, 0.4, "[Nombre(s) del grupo]", size=16, bold=True, color=WHITE)
text(s, 0.55, 5.55, 3.8, 0.4, "Optimización combinatoria · 2026", size=12, color=RGBColor(0xB5, 0xC4, 0xD4))

text(s, 5.2, 0.42, 7.9, 0.5, "La decisión puntual que toma el modelo", size=23, bold=True, color=NAVY)
q = box(s, 5.2, 1.08, 7.65, 1.15, fill=PALE_BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 0, 0, 0, 0, [[("Dado el estado actual (qué armas ya se asignaron y a qué objetivos), ", {}),
                      ("¿qué arma asigno ahora y a qué objetivo?", {"bold": True, "color": NAVY})]],
     size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=q)
text(s, 5.2, 2.42, 7.8, 0.75,
     [[("W armas, T objetivos. El objetivo j vale ", {}), ("Vⱼ", {"bold": True}), ("; el arma i lo destruye con prob. ", {}),
       ("pᵢⱼ", {"bold": True}), (". Cada arma va a ≤ 1 objetivo; un objetivo puede recibir varias.", {})]], size=13.5)
fpath = formula_png(r"$\min_x\; Z=\sum_{j=1}^{T} V_j \prod_{i=1}^{W}(1-p_{ij})^{x_{ij}}$", os.path.join(tmp, "f1.png"))
box(s, 5.2, 3.15, 7.65, 0.85, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
image(s, fpath, 5.35, 3.2, 4.4, 0.75, align="left")
text(s, 9.85, 3.2, 2.9, 0.75, "valor esperado que sobrevive (se minimiza)", size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

text(s, 5.2, 4.2, 7.8, 0.4, "Tarea 4 (MLP) → Tarea 5 (atención + puntero)", size=15, bold=True, color=NAVY)
rows = [("Pesos por posición", "pesos compartidos: cada par se evalúa con la misma función"),
        ("Solo W = 10, T = 8", "ningún peso depende de W ni de T (probado hasta 20 × 15)"),
        ("Orden fijo de armas", "la acción es un par (arma, objetivo): el modelo elige también el arma")]
for k, (old, new) in enumerate(rows):
    yy = 4.7 + k * 0.72
    chip(s, 5.2, yy, 2.55, 0.56, old, GRAY, size=12)
    arrow(s, 7.82, yy + 0.28, 8.22, yy + 0.28, color=MUTED)
    b_ = box(s, 8.3, yy, 4.55, 0.56, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    text(s, 0, 0, 0, 0, new, size=12, color=INK, anchor=MSO_ANCHOR.MIDDLE, shape=b_)

# ================================================================ 2. Descomposición
s = prs.slides.add_slide(BLANK)
header(s, 2, "De la solución del experto a decisiones",
       "Tuplas (estado, acciones posibles → acción escogida); cada paso agrega un par (arma libre → objetivo). Ejemplo real 10 × 8")
image(s, os.path.join(FIG, "descomposicion.png"), 0.35, 1.3, 12.65, 3.05)
steps = [("Estado inicial", "asignación vacía: las 10 armas libres (análogo a visited=[0])", NAVY),
         ("Acciones posibles", "todos los pares (arma libre, objetivo): k·T\n80 en el paso 0 … 8 en el último", BLUE),
         ("Acción escogida", "par del experto con mayor ganancia Vⱼ·qⱼ·pᵢⱼ (criterio MMR) → env.state_transition", RED),
         ("Término", "cuando las W armas están asignadas: W = 10 tuplas por instancia", GREEN)]
bw, gx, y0 = 2.95, 0.22, 4.5
for k, (title, body, col) in enumerate(steps):
    x = 0.5 + k * (bw + gx)
    card(s, x, y0, bw, 1.35, title, body, accent=col, title_size=13.5, body_size=11)
n_diff = sum(1 for st in D["steps"] if st["mmr"] != st["chosen"])
box(s, 0.5, 6.0, SW - 1.0, 0.85, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
text(s, 0.7, 6.02, SW - 1.4, 0.8,
     [[("¿Por qué ese orden? ", {"bold": True}),
       ("El experto entrega un conjunto de pares sin orden; ordenarlos por ganancia marginal da una etiqueta única por estado, "
        "reconstruye exactamente su solución y es coherente con MMR. En este ejemplo el par del experto difiere del que "
        f"elegiría MMR en {n_diff} de 10 pasos: eso es lo que el modelo debe aprender a corregir.", {})]],
     size=12, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 3. Matriz de entrada
s = prs.slides.add_slide(BLANK)
header(s, 3, f"Entrada: matriz (1 + W·T) × 16 — una fila por acción posible",
       f"{R['n_rows']} × 16 en entrenamiento. Todas las filas se procesan con los mismos pesos; ninguna columna depende del índice")
image(s, os.path.join(FIG, "estado.png"), 0.35, 1.3, 8.3, 2.7, align="left")
# esquema de filas
box(s, 8.85, 1.35, 4.0, 2.6, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 9.0, 1.42, 3.8, 0.35, "Filas", size=14, bold=True, color=NAVY)
chip(s, 9.0, 1.85, 1.2, 0.42, "fila 0", NAVY, size=11)
text(s, 10.3, 1.82, 2.5, 0.5, [[("global [CLS]", {"bold": True}), (": forma la query; siempre enmascarada", {})]], size=10.5)
chip(s, 9.0, 2.5, 1.2, 0.42, "1+i·T+j", BLUE, size=11)
text(s, 10.3, 2.47, 2.5, 0.6, [[("par (arma i → objetivo j)", {"bold": True}), (": estado + acción “asignar i a j”", {})]],
     size=10.5)
text(s, 9.0, 3.2, 3.8, 0.7, [[("Con k armas libres quedan ", {}), ("k·T filas válidas", {"bold": True}),
                             (" = las acciones de env.gen_actions.", {})]], size=10.5, color=MUTED)
cards = [
    ("Propias del par (0–2)", [[("0 pᵢⱼ", {"bold": True}), (" → efectividad del arma", {})],
                               [("1 Vⱼ / max V", {"bold": True}), (" → importancia del objetivo", {})],
                               [("2 pᵢⱼ − máxⱼ′ pᵢⱼ′", {"bold": True}), (" → ¿es su mejor destino?", {})]], BLUE, PALE_BLUE),
    ("Flags de estado (3–5)", [[("3 es_global", {"bold": True}), (" → fila de la query", {})],
                               [("4 no_disponible", {"bold": True, "color": RED}), (" → MÁSCARA (arma ya asignada)", {})],
                               [("5 asignado", {"bold": True}), (" → xᵢⱼ = 1 en la solución parcial", {})]], RED, PALE_RED),
    ("Relativas al estado (6–15)", [[("de las heurísticas: ", {}), ("ganancia Vⱼqⱼpᵢⱼ (MMR)", {"bold": True}),
                                     (", es_mmr, qⱼ, armas en j, ventaja y arrepentimiento vs otras armas libres, ", {}),
                                     ("completar con MMR", {"bold": True}), (" (anticipación), fracción restante", {})]],
     GREEN, PALE_GREEN),
]
cw = (SW - 1.0 - 2 * 0.2) / 3
for k, (title, body, col, fill) in enumerate(cards):
    card(s, 0.5 + k * (cw + 0.2), 4.2, cw, 1.75, title, body, accent=col, title_size=13.5, body_size=11.5, fill=fill)
box(s, 0.5, 6.1, SW - 1.0, 0.78, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
text(s, 0.7, 6.12, SW - 1.4, 0.74,
     [[("Independencia del tamaño: ", {"bold": True}),
       ("todo está normalizado por instancia (÷ máximo o fracciones) y la información global se repite en cada fila y en "
        "la fila 0. La self-attention deja que el par (i, j) compare con los pares (i′, j) de otras armas.", {})]],
     size=12, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 4. Salida y modelo
s = prs.slides.add_slide(BLANK)
header(s, 4, "Salida: un logit por fila, máscara y etiqueta",
       "Modelo: estructura de TSP_Attention; cambian VEC_LEN, la query de decisión (paso 3) y la columna de máscara (paso 4)")
blocks = [("Entrada", f"(1+W·T) × 16", NAVY, 1.35),
          ("Embedding", "Linear(16, 64)\ncompartido", BLUE, 1.45),
          ("Encoder × 2", "self-attention + FF\nresidual · LayerNorm\nDropout", BLUE, 1.75),
          ("Query q", "Wc[h_global ;\npromedio disp.]", PURPLE, 1.6),
          ("Puntero", "uᵣ = q·hᵣ / √d\n(Wq, Wk)", ORANGE, 1.45),
          ("Máscara", "col. 4 = 1\n→ uᵣ = −10⁹", RED, 1.35),
          ("Softmax", "P(fila r)", GREEN, 1.2)]
total = sum(w for *_, w in blocks)
gx = (SW - 1.0 - total) / (len(blocks) - 1)
x = 0.5
for k, (name, detail, col, w) in enumerate(blocks):
    b_ = box(s, x, 1.45, w, 1.25, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 0, 0, 0, 0, [[(name, {"bold": True, "size": 13})], [(detail, {"size": 10.5})]], color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b_)
    if k < len(blocks) - 1:
        arrow(s, x + w + 0.02, 2.07, x + w + gx - 0.02, 2.07, color=MUTED)
    x += w + gx
dp_final = R["final_config"]["dropout"]
text(s, 0.5, 2.85, 12.3, 0.4, f"tamaño moderado: d = 64 · 2 capas · {miles(R['final_params'])} parámetros · "
     f"Dropout evaluado en 0 / 0,1 / 0,2 (elegido por gap de validación: {num(dp_final, 1)})",
     size=11.5, color=MUTED, align=PP_ALIGN.CENTER)

image(s, os.path.join(FIG, "salida.png"), 0.4, 3.3, 4.6, 3.6)
card(s, 5.2, 3.35, 3.75, 3.5, "SALIDA Y ETIQUETA",
     [[("Logits: ", {"bold": True}), ("vector de largo 1 + W·T (uno por fila).", {"space_after": 4})],
      [("Acción ↔ fila: ", {"bold": True}), ("(\"constructive\", (i, j)) = fila 1 + i·T + j.", {"space_after": 4})],
      [("Máscara: ", {"bold": True}), ("fila global y pares de armas ya asignadas (no_disponible = 1) → prob. 0.",
                                       {"space_after": 4})],
      [("Y: ", {"bold": True}), ("índice entero de la fila del experto (≡ one-hot).", {"space_after": 4})],
      [("Pérdida: ", {"bold": True}), ("CrossEntropyLoss(model(x), Y).", {})]],
     accent=ORANGE, body_size=11)
rows = [("", "TSP_Attention", "WTAP_Attention"),
        ("VEC_LEN", "6", "16"),
        ("Query (paso 3)", "[h_actual ; h_inicial]\ncol. 2 + fila 0", "[h_global ; promedio disp.]\ncol. 3 (es_global)"),
        ("Máscara (paso 4)", "visitada (col. 3)", "no_disponible (col. 4)"),
        ("Encoder", "1 bloque", "2 bloques + LayerNorm + Dropout")]
tbl = s.shapes.add_table(len(rows), 3, Inches(9.15), Inches(3.35), Inches(3.7), Inches(3.5)).table
for c_, w in enumerate([1.05, 1.3, 1.35]):
    tbl.columns[c_].width = Inches(w)
for r_, row in enumerate(rows):
    tbl.rows[r_].height = Inches(0.5 if r_ in (0, 1) else 0.83)
    for c_, val in enumerate(row):
        cell = tbl.cell(r_, c_)
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r_ == 0 else (PALE_BLUE if c_ == 2 else (LIGHT if r_ % 2 else WHITE))
        text(s, 0, 0, 0, 0, val, size=9.5, bold=(r_ == 0 or c_ == 0), color=WHITE if r_ == 0 else INK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=cell)

# ================================================================ 5. Resultados
s = prs.slides.add_slide(BLANK)
header(s, 5, "Resultados: entrenamiento y evaluación (pasos opcionales 7 y 8)",
       f"{miles(R['n_train_samples'])} pares (estado, fila) de {miles(R['n_train_instances'])} instancias · " +
       "prueba en las mismas 50 instancias de la Tarea 4 · mismo GreedyAgent")
kp = [(pct(100 * R["val_acc_final"]), "accuracy validación", BLUE, "decisiones iguales al experto"),
      (pct(T_[MOD]["mean_gap"], 2), "gap modelo", BLUE, f"óptimo en {pct(T_[MOD]['pct_optimal'], 0)} de las instancias"),
      (pct(R["gap_mlp_t4"], 2), "gap MLP (Tarea 4)", GRAY, "orden fijo de armas"),
      (pct(T_[MMR]["mean_gap"], 2), "gap MMR", ORANGE, "heurística clásica"),
      (f"{pct(100 * R['model_vs_mmr']['wins'], 0)}", "gana a MMR", GREEN,
       f"empata en {pct(100 * R['model_vs_mmr']['ties'], 0)}")]
kw = (SW - 1.0 - 4 * 0.18) / 5
for k, (v, l, col, sub) in enumerate(kp):
    kpi(s, 0.5 + k * (kw + 0.18), 1.35, kw, 1.25, v, l, col, sub)
image(s, os.path.join(FIG, "curvas.png"), 0.4, 2.8, 6.2, 2.1, align="left")
image(s, os.path.join(FIG, "generalizacion.png"), 6.75, 2.75, 3.3, 2.2)
gen = R["generalization"]
card(s, 10.2, 2.8, 2.65, 2.15, "GENERALIZA",
     [[("Entrenado solo con 10×8, se usa sin cambios en ", {}), ("15×10 y 20×15", {"bold": True}),
       (" (151 y 301 filas).", {})]], accent=PURPLE, body_size=11)
g = {k: v for k, v in gen.items()}
d_mlp = T_[MOD]["mean_gap"] - R["gap_mlp_t4"]
E_ = R["experiments"]
dropout_line = ("gap de validación " + " · ".join(f"{num(v['config']['dropout'], 1)} → {pct(v['gap_val'], 2)}"
                                                  for v in E_.values())
                + (". Sin Dropout fue lo mejor: la pérdida de entrenamiento ≈ validación, el modelo no llega a sobreajustar."
                   if R["final_config"]["dropout"] == 0 else "."))
vs_mlp = "mejora a" if d_mlp < -0.25 else ("iguala a" if abs(d_mlp) <= 0.25 else "queda cerca de")
lines = " · ".join(f"{k}: {pct(v[MOD])} (MMR {pct(v[MMR])})" for k, v in g.items())
box(s, 0.5, 5.1, SW - 1.0, 1.75, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 0.7, 5.15, SW - 1.4, 1.7,
     [[("Lectura. ", {"bold": True}),
       (f"Con la misma receta de imitación, el puntero sobre pares {vs_mlp} la MLP de la Tarea 4 "
        f"({pct(T_[MOD]['mean_gap'], 2)} vs {pct(R['gap_mlp_t4'], 2)}) y reduce el gap de MMR "
        f"({pct(T_[MMR]['mean_gap'], 2)}), pero ahora decide también qué arma usar, con pesos compartidos y para "
        "cualquier tamaño.", {"space_after": 5})],
      [("Gap vs experto por tamaño: ", {"bold": True}), (lines + ".", {"space_after": 5})],
      [("Dropout: ", {"bold": True}), (dropout_line, {"space_after": 5})],
      [("Próximo paso: ", {"bold": True}), ("el correo lo deja para después: RL o DAgger para corregir el distribution "
                                            "shift y más capas/cabezas.", {})]],
     size=12)

prs.save(OUT)
for f in os.listdir(tmp):
    os.remove(os.path.join(tmp, f))
os.rmdir(tmp)
print("Presentación guardada en", OUT)
