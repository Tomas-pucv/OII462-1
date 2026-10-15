"""Genera la presentación final del módulo (WTAP con Transformer + puntero) a partir de figuras_proyecto/.

Uso (desde la raíz del repositorio):
    python scripts/make_presentation_proyecto.py
    soffice --headless --convert-to pdf Proyecto_WTAP_presentacion.pptx
"""
import json
import os
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figuras_proyecto")
OUT = os.path.join(ROOT, "Proyecto_WTAP_presentacion.pptx")
R = json.load(open(os.path.join(FIG, "resultados.json"), encoding="utf-8"))

FONT = "Arial"
NAVY = RGBColor(0x1B, 0x2A, 0x41)
INK = RGBColor(0x2E, 0x35, 0x40)
MUTED = RGBColor(0x6B, 0x75, 0x85)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
ORANGE = RGBColor(0xE8, 0x71, 0x0A)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
PURPLE = RGBColor(0x94, 0x67, 0xBD)
RED = RGBColor(0xC0, 0x39, 0x2B)
PALE_BLUE = RGBColor(0xDC, 0xEA, 0xF7)
PALE_ORANGE = RGBColor(0xFD, 0xEB, 0xD9)
PALE_RED = RGBColor(0xF9, 0xDE, 0xDC)
PALE_GREEN = RGBColor(0xDF, 0xF0, 0xDF)
PALE_PURPLE = RGBColor(0xEC, 0xE4, 0xF5)

SECTIONS = {1: ("1 · PROBLEMA", BLUE), 2: ("2 · REPRESENTACIÓN", GREEN), 3: ("3 · DATOS", PURPLE),
            4: ("4 · ARQUITECTURA", ORANGE), 5: ("5 · EXPERIMENTOS", RED)}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]
SLIDES = []          # (slide, notes)


# ---------------------------------------------------------------- helpers (base: presentación de la Tarea 5)
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, line_w=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, shape=None):
    """runs: str, o lista de párrafos; cada párrafo es str o lista de (texto, {opciones})."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = runs if isinstance(runs, list) else [runs]
    for k, par in enumerate(paragraphs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        pieces = par if isinstance(par, list) else [(par, {})]
        for txt, opt in pieces:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = opt.get("font", FONT)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
        if isinstance(par, list) and par and "space_after" in par[0][1]:
            p.space_after = Pt(par[0][1]["space_after"])
    return shape


def image(slide, path, x, y, w=None, h=None, align="center"):
    """Inserta la imagen ajustada a la caja (x, y, w, h) manteniendo la proporción."""
    iw, ih = Image.open(path).size
    ratio = iw / ih
    if w is not None and h is not None:
        if w / h > ratio:
            nw, nh = h * ratio, h
        else:
            nw, nh = w, w / ratio
        ox = x + {"center": (w - nw) / 2, "left": 0, "right": w - nw}[align]
        oy = y + (h - nh) / 2
    elif w is not None:
        nw, nh, ox, oy = w, w / ratio, x, y
    else:
        nw, nh, ox, oy = h * ratio, h, x, y
    return slide.shapes.add_picture(path, Inches(ox), Inches(oy), Inches(nw), Inches(nh))


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def new_slide(section, title, subtitle=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    n = len(prs.slides)
    label, col = SECTIONS[section]
    box(s, 0, 0, SW, SH, fill=WHITE)
    box(s, 0.5, 0.42, 0.09, 0.62, fill=col)
    text(s, 0.72, 0.3, 10.3, 0.6, title, size=26, bold=True, color=NAVY)
    if subtitle:
        text(s, 0.74, 0.86, 11.9, 0.4, subtitle, size=13.5, color=MUTED)
    chip(s, SW - 2.75, 0.4, 2.25, 0.36, label, col, size=10.5)
    text(s, 0.5, 7.07, 9, 0.3, "OII462 · Proyecto del módulo · Transformer con puntero para el WTAP", size=10, color=MUTED)
    text(s, SW - 1.4, 7.07, 0.9, 0.3, f"{n} / {TOTAL}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    SLIDES.append((s, notes))
    return s


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=14, body_size=12, fill=LIGHT):
    box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(slide, x, y + 0.15, 0.07, h - 0.3, fill=accent)
    text(slide, x + 0.22, y + 0.1, w - 0.35, 0.4, title, size=title_size, bold=True, color=accent)
    text(slide, x + 0.22, y + 0.48, w - 0.35, h - 0.55, body, size=body_size, color=INK, spacing=1.1)


def kpi(slide, x, y, w, h, value, label, color, sub=None, vsize=24):
    box(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(slide, x, y, w, 0.08, fill=color)
    text(slide, x, y + 0.12, w, 0.55, value, size=vsize, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(slide, x + 0.05, y + 0.64, w - 0.1, 0.35, label, size=11, bold=True, color=INK, align=PP_ALIGN.CENTER)
    if sub:
        text(slide, x + 0.05, y + 0.92, w - 0.1, 0.3, sub, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def formula_png(latex, name, fontsize=26, color="#1B2A41"):
    path = os.path.join(TMP, f"{name}.png")
    fig = plt.figure(figsize=(8, 1))
    fig.text(0, 0.5, latex, fontsize=fontsize, color=color, va="center")
    fig.savefig(path, dpi=250, bbox_inches="tight", pad_inches=0.05, transparent=True)
    plt.close(fig)
    return path


def chip(slide, x, y, w, h, label, fill, size=12, color=WHITE, bold=True):
    s = box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    text(slide, 0, 0, 0, 0, label, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=s)
    return s


def block(slide, x, y, w, h, title, detail, fill, size_t=12.5, size_d=10.5, color=WHITE):
    b = box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    pars = [[(title, {"bold": True, "size": size_t})]]
    if detail:
        pars.append([(detail, {"size": size_d})])
    text(slide, 0, 0, 0, 0, pars, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b, spacing=1.0)
    return b


def table(slide, x, y, w, rows, widths, row_h=0.36, size=10.5, header_fill=NAVY, fills=None, bold_first_col=False,
          aligns=None, header_size=None):
    t = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for c_, wd in enumerate(widths):
        t.columns[c_].width = Inches(wd)
    for r_, row in enumerate(rows):
        t.rows[r_].height = Inches(row_h)
        for c_, val in enumerate(row):
            cell = t.cell(r_, c_)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            if r_ == 0:
                cell.fill.fore_color.rgb = header_fill
            elif fills and fills[r_ - 1] is not None:
                cell.fill.fore_color.rgb = fills[r_ - 1]
            else:
                cell.fill.fore_color.rgb = LIGHT if r_ % 2 else WHITE
            al = aligns[c_] if aligns else PP_ALIGN.LEFT
            text(slide, 0, 0, 0, 0, val, size=(header_size or size) if r_ == 0 else size,
                 bold=(r_ == 0 or (bold_first_col and c_ == 0)), color=WHITE if r_ == 0 else INK,
                 align=al, anchor=MSO_ANCHOR.MIDDLE, shape=cell, spacing=1.0)
    return t


def pct(v, d=1):
    return f"{v:.{d}f}%".replace(".", ",")


def num(v, d=2):
    return f"{v:.{d}f}".replace(".", ",")


def miles(n):
    return f"{n:,}".replace(",", ".")


def pm(m, s, d=2, unit=""):
    return f"{num(m, d)} ± {num(s, d)}{unit}"


TOTAL = 14
TMP = os.path.join(FIG, "_tmp")
os.makedirs(TMP, exist_ok=True)
T_ = R["test"]
MOD, MMR, RND, EXP = "Modelo (atención)", "MMR (heurística clásica)", "Aleatorio", "Experto (ILS)"
D = R["demo"]
HP = R["hparams"]
CV = R["convergence"]
SZ = R["sizes"]
TM = R["times_ms"]
size_labels = list(SZ)

# ================================================================ 1. Portada + agenda
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=WHITE)
box(s, 0, 0, 5.0, SH, fill=NAVY)
text(s, 0.55, 0.7, 4.2, 0.4, "OII462 · PROYECTO DEL MÓDULO", size=13, bold=True, color=RGBColor(0x8F, 0xB8, 0xDE))
text(s, 0.55, 1.15, 4.3, 2.6, "Transformer con puntero para el Weapon-Target Assignment Problem", size=30, bold=True,
     color=WHITE, spacing=1.0)
text(s, 0.55, 3.6, 4.1, 1.4, "Un encoder–decoder con atención que aprende a asignar armas a objetivos imitando paso a "
     "paso a un algoritmo experto", size=14.5, color=RGBColor(0xD6, 0xE2, 0xEE), spacing=1.15)
box(s, 0.55, 5.1, 0.9, 0.05, fill=ORANGE)
text(s, 0.55, 5.3, 4.1, 0.4, "[Nombre(s) del grupo]", size=16, bold=True, color=WHITE)
text(s, 0.55, 5.7, 4.1, 0.4, "Optimización combinatoria · 2026", size=12, color=RGBColor(0xB5, 0xC4, 0xD4))
text(s, 5.6, 0.6, 7.2, 0.5, "Agenda (13–15 min)", size=22, bold=True, color=NAVY)
agenda = [("1", "Introducción al problema", "definición formal y ejemplo visual", "1,5 min", BLUE),
          ("2", "Representación del estado", "secuencia de entrada y justificación de cada feature", "2 min", GREEN),
          ("3", "Generación de datos", "experto, descomposición y codificación de la salida", "2 min", PURPLE),
          ("4", "Arquitectura", "encoder + decoder con puntero", "3 min", ORANGE),
          ("5", "Experimentos y resultados", "configuración, convergencia, evaluación, tamaños, conclusiones", "4,5 min", RED)]
for k, (n, t, d, m, col) in enumerate(agenda):
    y = 1.35 + k * 1.05
    c = box(s, 5.6, y, 0.7, 0.7, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, 0, 0, 0, 0, n, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
    text(s, 6.5, y - 0.02, 5.0, 0.45, t, size=17, bold=True, color=NAVY)
    text(s, 6.5, y + 0.37, 5.2, 0.4, d, size=12, color=MUTED)
    text(s, 11.6, y + 0.12, 1.2, 0.45, m, size=13, bold=True, color=col, align=PP_ALIGN.RIGHT)
SLIDES.append((s, "Portada (≈15 s). Presentamos el modelo encoder–decoder que construimos para el WTAP, la agenda "
                  "y los tiempos de cada sección."))

# ================================================================ 2. Problema: definición formal
s = new_slide(1, "Weapon-Target Assignment: definición formal",
              "Asignar W armas a T objetivos para minimizar el valor esperado que sobrevive al ataque",
              notes="(≈45 s) El WTAP: tenemos W armas y T objetivos. Cada objetivo tiene un valor V_j y cada arma i lo "
                    "destruye con probabilidad p_ij. La salida es una matriz binaria X; cada arma va a lo más a un "
                    "objetivo, pero un objetivo puede recibir varias armas. Minimizamos el valor esperado que sobrevive: "
                    "es no lineal por el producto de probabilidades y el problema es NP-hard.")
items = [("ENTRADAS", BLUE, [[("W armas, T objetivos", {"bold": True})],
                             [("Vⱼ > 0: valor (amenaza) del objetivo j", {})],
                             [("pᵢⱼ ∈ [0, 1]: prob. de que el arma i destruya j (qᵢⱼ = 1 − pᵢⱼ sobrevive)", {})]]),
         ("SALIDA", GREEN, [[("X = [xᵢⱼ] binaria (W × T): ", {"bold": True}), ("xᵢⱼ = 1 si el arma i va al objetivo j", {})]]),
         ("RESTRICCIÓN", RED, [[("Σⱼ xᵢⱼ ≤ 1  ∀ i", {"bold": True}), (": cada arma a lo más a un objetivo. ", {}),
                                ("Un objetivo puede recibir varias armas.", {})]])]
y = 1.45
for title, col, body in items:
    hgt = 1.45 if title == "ENTRADAS" else 0.95
    card(s, 0.5, y, 6.3, hgt, title, body, accent=col, title_size=13, body_size=12)
    y += hgt + 0.15
fpath = formula_png(r"$\min_x\; Z=\sum_{j=1}^{T} V_j \prod_{i=1}^{W}(1-p_{ij})^{x_{ij}}$", "obj")
card(s, 0.5, y, 6.3, 1.45, "FUNCIÓN OBJETIVO", "", accent=PURPLE, title_size=13)
image(s, fpath, 0.8, y + 0.45, 4.1, 0.85, align="left")
text(s, 4.95, y + 0.45, 1.75, 0.9, "valor esperado que sobrevive", size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
image(s, os.path.join(ROOT, "figuras", "instancia.png"), 7.05, 1.45, 5.8, 3.4)
facts = [("NP-hard", "objetivo no lineal (Lloyd & Witsenhausen, 1986)"), ("Tᵂ = 8¹⁰ ≈ 10⁹", "asignaciones en 10 × 8"),
         ("Vⱼ ~ U{25..100}", "pᵢⱼ ~ U(0,6; 0,9) (Ahuja et al., 2007)")]
for k, (big, small) in enumerate(facts):
    x = 7.05 + k * 1.97
    box(s, x, 5.0, 1.87, 1.3, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, x, 5.07, 1.87, 0.45, big, size=15, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x + 0.08, 5.52, 1.71, 0.75, small, size=10, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 7.05, 6.4, 5.8, 0.5, "Instancia de ejemplo 10 × 8: probabilidades pᵢⱼ (izq.) y valores Vⱼ (der.)",
     size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 3. Ejemplo visual: estado → acción
s = new_slide(1, "Un estado y su acción", "El modelo construye la solución par a par: en cada paso elige qué arma asignar y a qué objetivo",
              notes="(≈45 s) Así se ve una decisión. El estado es la asignación parcial: dos armas ya asignadas (azul). "
                    "Las acciones posibles son todos los pares arma libre–objetivo (líneas grises, 64 en este paso). "
                    "La acción es el par que escoge el experto (rojo). Después se actualiza la probabilidad de "
                    "supervivencia del objetivo y se repite hasta asignar las 10 armas.")
image(s, os.path.join(FIG, "instancia_estado_accion.png"), 0.35, 1.3, 12.65, 3.75)
ec = D["ej_chosen"]
cards3 = [("ESTADO", BLUE, [[(f"Asignación parcial: {10 - (D['ej_n_actions'] // 8)} armas ya asignadas (azul). ", {}),
                            ("qⱼ = prob. de que j sobreviva con lo asignado (color del objetivo).", {})]]),
          ("ACCIONES POSIBLES", GRAY, [[("Todos los pares (arma libre, objetivo): ", {}),
                                       (f"{D['ej_n_actions']} en este paso", {"bold": True}),
                                       (" (80 al inicio, 8 al final).", {})]]),
          ("ACCIÓN ESCOGIDA", RED, [[(f"Arma {ec[0]} → objetivo {ec[1]} (rojo). ", {"bold": True}),
                                    ("Transición: xᵢⱼ = 1 y qⱼ ← qⱼ · (1 − pᵢⱼ).", {})]])]
cw = (SW - 1.0 - 2 * 0.2) / 3
for k, (t, col, body) in enumerate(cards3):
    card(s, 0.5 + k * (cw + 0.2), 5.25, cw, 1.6, t, body, accent=col, title_size=13, body_size=12)

# ================================================================ 4. Secuencia de entrada
s = new_slide(2, "Secuencia de entrada: una fila por acción posible",
              f"Matriz (1 + W·T) × 16 — {R['n_rows']} × 16 en 10 × 8. Todas las filas pasan por los mismos pesos",
              notes="(≈40 s) Cada elemento de la secuencia es una ACCIÓN: el par arma i → objetivo j, en la fila 1 + i·T + j. "
                    "Elegimos pares porque es exactamente lo que evalúa la heurística clásica MMR y porque así la máscara es "
                    "natural. La fila 0 es una fila global, como el token CLS, que resume el estado. Cada fila tiene 16 "
                    "features en tres roles: propias del par, flags y relativas al estado.")
image(s, os.path.join(FIG, "estado.png"), 0.35, 1.35, 8.4, 2.75, align="left")
box(s, 8.95, 1.38, 3.9, 2.7, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 9.1, 1.45, 3.6, 0.35, "¿Qué es cada elemento?", size=14, bold=True, color=NAVY)
chip(s, 9.1, 1.9, 1.15, 0.42, "fila 0", NAVY, size=11)
text(s, 10.35, 1.86, 2.45, 0.6, [[("fila global [CLS]", {"bold": True}), (": resume el estado; forma la query", {})]],
     size=10.5)
chip(s, 9.1, 2.6, 1.15, 0.42, "1 + i·T + j", BLUE, size=10.5)
text(s, 10.35, 2.56, 2.45, 0.7, [[("par (arma i → objetivo j)", {"bold": True}), (": el estado visto desde la acción "
                                                                                 "“asignar i a j”", {})]], size=10.5)
text(s, 9.1, 3.35, 3.65, 0.7, [[("Con k armas libres hay ", {}), ("k·T filas válidas", {"bold": True}),
                               (" = las acciones del entorno.", {})]], size=10.5, color=MUTED)
roles = [("Propias del par (0–2)", "describen la acción por sí misma: efectividad pᵢⱼ y valor Vⱼ", BLUE, PALE_BLUE),
         ("Flags de estado (3–5)", "query (es_global), máscara (no_disponible) y pares ya asignados", RED, PALE_RED),
         ("Relativas al estado (6–15)", "lo que miran las heurísticas: ganancia marginal (MMR), regret y anticipación",
          GREEN, PALE_GREEN)]
for k, (t, body, col, fill) in enumerate(roles):
    card(s, 0.5 + k * (cw + 0.2), 4.3, cw, 1.3, t, body, accent=col, title_size=13, body_size=11.5, fill=fill)
box(s, 0.5, 5.75, SW - 1.0, 1.1, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 0.7, 5.78, SW - 1.4, 1.05,
     [[("¿Por qué pares y no objetivos? ", {"bold": True}),
       ("MMR, la heurística constructiva clásica del WTAP, evalúa la ganancia de cada par (arma libre, objetivo). Con una fila "
        "por par, el puntero elige arma y objetivo a la vez (sin orden fijo de armas), la máscara es natural y el tamaño de la "
        "secuencia se adapta a cualquier W y T. Todo está normalizado por instancia: ningún peso depende del índice de la fila.",
        {})]], size=11.5, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 5–6. Features
FEATS = [
    ("0", "p_ij", "prob. de que el arma i destruya j", "define la efectividad de la asignación",
     "pares con p alto reducen más el costo"),
    ("1", "valor", "Vⱼ / max V", "el costo pondera cada objetivo por Vⱼ", "conviene cubrir antes los objetivos valiosos"),
    ("2", "p_vs_mejor", "pᵢⱼ − maxⱼ′ pᵢⱼ′ (≤ 0)", "costo de oportunidad del arma i",
     "0 ⇒ j es el mejor uso posible de esa arma"),
    ("3", "es_global", "1 solo en la fila 0", "el decoder debe ubicar la fila que resume el estado",
     "forma la query del puntero"),
    ("4", "no_disponible", "1 si el arma i ya fue asignada (o fila 0)", "codifica la restricción Σⱼ xᵢⱼ ≤ 1",
     "máscara: logit −10⁹ ⇒ probabilidad 0"),
    ("5", "asignado", "1 si xᵢⱼ = 1 en la solución parcial", "contexto: qué pares ya se eligieron",
     "la atención ve qué objetivos ya están cubiertos"),
    ("6", "q_sobrevive", "qⱼ con lo asignado", "rendimientos decrecientes: cada arma actúa sobre lo que queda",
     "qⱼ bajo ⇒ otra arma en j gana poco"),
    ("7", "valor_en_juego", "Vⱼ · qⱼ / max V", "es el término del costo asociado a j",
     "mucho valor en juego atrae armas"),
    ("8", "ganancia ★", "Vⱼ qⱼ pᵢⱼ / máx. de los pares válidos", "reducción inmediata del costo: el criterio de MMR",
     "candidato natural; el modelo aprende cuándo no seguirlo"),
    ("9", "es_mmr", "1 en el par que elegiría MMR", "resume la decisión greedy en una columna",
     "referencia explícita que el modelo corrige"),
    ("10", "armas_en_obj", "nº de armas ya en j / W", "saturación del objetivo", "evita concentrar demasiadas armas"),
    ("11", "ventaja", "pᵢⱼ − media de p de las armas libres en j", "¿es i comparativamente buena para j?",
     "favorece usar el arma “especialista”"),
    ("12", "arrepentimiento", "gᵢⱼ − mejor g de otra arma libre en j", "regret: ¿otra arma lo haría mejor?",
     "negativo ⇒ conviene reservar j para otra arma"),
    ("13", "mmr_completa ★", "1 si al completar con MMR el arma i va a j", "anticipa las decisiones futuras",
     "en la Tarea 4, la anticipación fue lo más valioso"),
    ("14", "q_final_mmr", "qⱼ al completar con MMR", "¿quedará j bien cubierto al final?",
     "evita sobre-asignar lo que se cubrirá después"),
    ("15", "frac_restantes", "armas libres / W (en todas las filas)", "etapa de la construcción",
     "al inicio hay más margen; al final hay que cerrar"),
]
ROLE_FILL = {"0": PALE_BLUE, "1": PALE_BLUE, "2": PALE_BLUE, "3": PALE_RED, "4": PALE_RED, "5": PALE_RED}
for part, (lo, hi) in enumerate([(0, 8), (8, 16)]):
    s = new_slide(2, f"Features por elemento ({lo}–{hi - 1}): qué capturan y por qué",
                  "Azul = propias del par · rojo = flags de estado · verde = relativas al estado (★ = las más importantes)",
                  notes=("(≈40 s) Primeras 8 features. Las tres primeras describen la acción por sí misma. Los flags le dicen "
                         "al modelo dónde está la fila de la query y qué acciones son inválidas. q_j y el valor en juego "
                         "capturan los rendimientos decrecientes: un segundo disparo solo actúa sobre lo que sobrevivió."
                         if part == 0 else
                         "(≈40 s) Las features relativas vienen de las heurísticas: la ganancia marginal es el criterio de "
                         "MMR; ventaja y arrepentimiento comparan con otras armas libres (regret); mmr_completa y q_final "
                         "miran hacia adelante completando la solución con MMR. Así el modelo puede aprender cuándo NO "
                         "conviene tomar la mayor ganancia inmediata."))
    rows = [("#", "Feature", "Qué captura", "Por qué es relevante", "Relación con la decisión")]
    fills = []
    for f in FEATS[lo:hi]:
        rows.append((f[0], [[(f[1], {"font": "Courier New", "size": 11, "bold": True})]], f[2], f[3], f[4]))
        fills.append(ROLE_FILL.get(f[0], PALE_GREEN))
    table(s, 0.5, 1.4, SW - 1.0, rows, [0.5, 2.0, 3.25, 3.3, 3.28], row_h=0.66, size=11.5, fills=fills,
          aligns=[PP_ALIGN.CENTER] + [PP_ALIGN.LEFT] * 4, header_size=12)

# ================================================================ 7. Generación de datos
s = new_slide(3, "Generación de datos: de la instancia al par (x, y)",
              "Aprendizaje por imitación: el experto resuelve, la solución se descompone en decisiones y cada decisión es un ejemplo",
              notes="(≈2 min) El experto es una Búsqueda Local Iterada: parte de MMR, aplica búsqueda local con movimientos "
                    "reassign y swap, y perturba. La elegimos porque alcanza el óptimo exacto en todas las instancias que "
                    "verificamos y tarda milisegundos. Su solución es un conjunto de pares sin orden: la descomponemos "
                    "agregando en cada paso el par con mayor ganancia marginal (el criterio de MMR). Cada paso da un par "
                    "(estado, acción): x es la matriz state2vec y y es el índice de la fila escogida, que equivale a un "
                    "one-hot sobre las filas. Con 6.000 instancias obtenemos 60.000 ejemplos.")
flow = [("Instancia", "Vⱼ ~ U{25..100}\npᵢⱼ ~ U(0,6; 0,9)\nW = 10, T = 8", NAVY),
        ("Experto: ILS", "MMR + búsqueda local\n(reassign, swap)\n+ perturbación", GREEN),
        ("Solución", "10 pares (i, a*ᵢ)\nsin orden", BLUE),
        ("Descomposición", "en cada paso: par del\nexperto con mayor\nganancia Vⱼqⱼpᵢⱼ", PURPLE),
        ("Ejemplo (x, y)", "x = state2vec (81 × 16)\ny = 1 + i·T + j\n(10 por instancia)", ORANGE)]
fw, gap_ = 2.2, (SW - 1.0 - 5 * 2.2) / 4
for k, (t, d, col) in enumerate(flow):
    x = 0.5 + k * (fw + gap_)
    block(s, x, 1.45, fw, 1.6, t, d, col, size_t=14, size_d=11)
    if k < 4:
        arrow(s, x + fw + 0.03, 2.25, x + fw + gap_ - 0.03, 2.25, color=MUTED, width=2.5)
exact_10 = "10/10 (10 × 8, meet-in-the-middle)"
exact_5 = R.get("exact_check_5x4", {}).get("5×4")
card(s, 0.5, 3.3, 4.1, 2.25, "¿POR QUÉ ESTE EXPERTO?",
     [[("Óptimo verificado: ", {"bold": True}), (f"{exact_10}" + (f" y {exact_5}/30 (5 × 4, fuerza bruta)." if exact_5 else "."),
                                                {"space_after": 4})],
      [("Rápido: ", {"bold": True}), (f"≈{num(TM[EXP], 0)} ms por instancia; el exacto tarda ≈2 s en 10 × 8 y no escala.",
                                      {"space_after": 4})],
      [("Constructivo vs búsqueda local: ", {"bold": True}),
       ("imitamos la solución final, no la trayectoria de movimientos (ligada a perturbaciones aleatorias).", {})]],
     accent=GREEN, body_size=11)
card(s, 4.75, 3.3, 4.0, 2.25, "DESCOMPOSICIÓN",
     [[("Estado inicial: ", {"bold": True}), ("ninguna arma asignada.", {"space_after": 3})],
      [("Acción: ", {"bold": True}), ("par (arma libre, objetivo); k·T posibles.", {"space_after": 3})],
      [("Orden: ", {"bold": True}), ("mayor ganancia marginal entre los pares del experto pendientes → etiqueta única por "
                                    "estado y reconstruye exactamente su solución.", {"space_after": 3})],
      [("Término: ", {"bold": True}), ("las W armas asignadas (10 decisiones).", {})]], accent=PURPLE, body_size=11)
card(s, 8.9, 3.3, 3.95, 2.25, "SALIDA (ETIQUETA)",
     [[("y = índice de la fila escogida", {"bold": True}), (" (entero) ≡ vector one-hot de largo 1 + W·T.", {"space_after": 4})],
      [("Las filas enmascaradas nunca son la etiqueta (verificado en todas las muestras).", {"space_after": 4})],
      [("Pérdida: ", {"bold": True}), ("CrossEntropy sobre los logits enmascarados.", {})]], accent=ORANGE, body_size=11)
ds = [(miles(R["n_train_instances"]), "instancias de entrenamiento", f"→ {miles(R['n_train_samples'])} pares (x, y)"),
      ("300", "instancias de validación", f"→ {miles(R['n_val_samples'])} pares"),
      (str(R["n_test_instances"]), "instancias de prueba (10 × 8)", "las mismas de la Tarea 4"),
      ("+ 120", "instancias por tamaño", "5×4, 15×10, 20×15, 30×20")]
for k, (v, l, sub) in enumerate(ds):
    kpi(s, 0.5 + k * 3.13, 5.7, 2.95, 1.2, v, l, [NAVY, BLUE, GREEN, PURPLE][k], sub, vsize=21)

# ================================================================ 8. Arquitectura (diagrama)
s = new_slide(4, "Arquitectura: encoder + decoder con puntero",
              "El encoder contextualiza la secuencia; el decoder convierte el estado en una distribución sobre las acciones",
              notes="(≈2 min) El encoder recibe la secuencia de 81 filas con sus 16 features. Un embedding lineal compartido "
                    "las lleva a dimensión 64 y dos capas de self-attention con feed-forward, residual y LayerNorm hacen que "
                    "cada par mire a todos los demás. Sale H: un vector de 64 por fila, que ya sabe del resto de la instancia "
                    "y del estado. El decoder toma H, la fila global y la máscara: arma un contexto con el embedding de la "
                    "fila global y el promedio de los pares disponibles, lo proyecta a una query q, y calcula la "
                    "compatibilidad de q con cada fila (puntero). Las filas no disponibles reciben −10⁹ y el softmax da la "
                    "probabilidad de cada acción. Se ejecuta la acción y, como las features cambian, el encoder se vuelve a "
                    "aplicar en el paso siguiente.")
# contenedores
box(s, 0.45, 1.35, 6.05, 5.55, fill=RGBColor(0xF1, 0xF6, 0xFC), line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    radius=0.04, line_w=2)
box(s, 6.85, 1.35, 6.05, 5.55, fill=RGBColor(0xFE, 0xF6, 0xEE), line=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    radius=0.04, line_w=2)
text(s, 0.65, 1.42, 3.5, 0.4, "ENCODER", size=17, bold=True, color=BLUE)
text(s, 7.05, 1.42, 3.5, 0.4, "DECODER", size=17, bold=True, color=ORANGE)
text(s, 2.2, 1.47, 4.2, 0.35, "2 capas · 1 cabeza · d = 64 · se recalcula en cada paso", size=10.5, color=MUTED,
     align=PP_ALIGN.RIGHT)
text(s, 8.6, 1.47, 4.2, 0.35, "1 capa de decisión (sin glimpse) · W pasos", size=10.5, color=MUTED, align=PP_ALIGN.RIGHT)


def tag(slide, x, y, label, col):
    text(slide, x, y, 0.95, 0.3, label, size=9, bold=True, color=col)


# encoder: columna
ex, ew = 1.55, 4.75
tag(s, 0.6, 2.05, "ENTRA", BLUE)
block(s, ex, 1.95, ew, 0.72, "Secuencia de entrada: (1 + W·T) × 16",
      "fila 0 = global [CLS] · filas 1 + i·T + j = pares (arma i, objetivo j)", NAVY, size_t=12, size_d=10)
arrow(s, ex + ew / 2, 2.69, ex + ew / 2, 2.86, color=MUTED)
tag(s, 0.6, 2.95, "HACE", BLUE)
block(s, ex, 2.88, ew, 0.55, "Embedding lineal compartido 16 → 64", "mismos pesos para todas las filas", BLUE,
      size_t=12, size_d=10)
arrow(s, ex + ew / 2, 3.45, ex + ew / 2, 3.62, color=MUTED)
box(s, ex - 0.08, 3.62, ew + 0.16, 1.95, fill=None, line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, ex + ew - 0.75, 3.64, 0.8, 0.3, "× 2", size=13, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)
block(s, ex + 0.1, 3.75, ew - 0.95, 0.72, "Self-attention (1 cabeza)",
      "A = softmax(QKᵀ/√d) · cada par mira a todos\n+ residual + LayerNorm", RGBColor(0x2E, 0x86, 0xC1), size_t=11.5,
      size_d=9.5)
arrow(s, ex + (ew - 0.85) / 2, 4.49, ex + (ew - 0.85) / 2, 4.65, color=MUTED)
block(s, ex + 0.1, 4.67, ew - 0.95, 0.75, "Feed-forward por fila",
      "64 → 128 → 64 (ReLU)\n+ residual + LayerNorm", RGBColor(0x2E, 0x86, 0xC1), size_t=11.5, size_d=9.5)
arrow(s, ex + ew / 2, 5.59, ex + ew / 2, 5.78, color=MUTED)
tag(s, 0.6, 5.9, "SALE", BLUE)
block(s, ex, 5.8, ew, 0.95, "H ∈ ℝ^((1+W·T) × 64)",
      "un embedding contextualizado por fila: cada par “sabe” del estado\nglobal y de los demás pares de la instancia",
      RGBColor(0x15, 0x4F, 0x7A), size_t=12.5, size_d=10)
# flecha encoder → decoder
arrow(s, ex + ew + 0.02, 6.27, 7.95, 6.27, color=NAVY, width=3)
text(s, 6.25, 6.3, 1.0, 0.3, "H", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
# decoder: columna (de abajo hacia arriba)
dx, dw = 7.95, 4.75
tag(s, 7.0, 5.9, "ENTRA", ORANGE)
block(s, dx, 5.8, dw, 0.95, "H + flag es_global (col. 3) + máscara (col. 4)",
      "estado parcial ya codificado en H · qué filas son acciones válidas", RGBColor(0x8A, 0x4B, 0x08), size_t=11.5,
      size_d=10)
arrow(s, dx + dw / 2, 5.78, dx + dw / 2, 5.6, color=MUTED)
tag(s, 7.0, 4.95, "HACE", ORANGE)
block(s, dx, 4.85, dw, 0.72, "Contexto: [h_global ; promedio de h disponibles]",
      "fila [CLS] + embedding del grafo (128) → q = W_c[·] ∈ ℝ⁶⁴", ORANGE, size_t=11.5, size_d=10)
arrow(s, dx + dw / 2, 4.83, dx + dw / 2, 4.66, color=MUTED)
block(s, dx, 3.93, dw, 0.7, "Puntero (atención de 1 consulta)", "uᵣ = (W_q q) · (W_k hᵣ) / √64  — un logit por fila", ORANGE,
      size_t=11.5, size_d=10)
arrow(s, dx + dw / 2, 3.91, dx + dw / 2, 3.74, color=MUTED)
block(s, dx, 3.06, dw, 0.65, "Máscara", "uᵣ = −10⁹ si no_disponible = 1 (arma usada o fila global)", RED, size_t=11.5,
      size_d=10)
arrow(s, dx + dw / 2, 3.04, dx + dw / 2, 2.87, color=MUTED)
tag(s, 7.0, 2.05, "SALE", ORANGE)
block(s, dx, 1.95, dw, 0.9, "Softmax → P(acción = fila r)",
      "distribución sobre las k·T acciones válidas\nGreedyAgent toma el argmax → env.state_transition", GREEN, size_t=12.5,
      size_d=10)

# ================================================================ 9. Arquitectura en detalle
s = new_slide(4, "Arquitectura en detalle y comparación con Glimpse",
              f"{miles(R['final_params'])} parámetros entrenables · ningún peso depende de W ni de T",
              notes="(≈1 min) En detalle: cada etapa con su forma. La query se arma con la fila global y el promedio de los "
                    "pares disponibles, como el embedding del grafo de Kool et al. El puntero es una atención con una sola "
                    "consulta: sus pesos no mezclan valores, son directamente la política. Respecto de la arquitectura "
                    "Glimpse vista en clases: Kool agrega antes del puntero un glimpse, una atención multi-cabeza del "
                    "contexto sobre H, y recorta los logits con tanh. Nuestro decoder no tiene glimpse; es la principal "
                    "mejora pendiente.")
rows = [("Etapa", "Operación", "Forma (10 × 8)"),
        ("Entrada", "state2vec(estado)", "81 × 16"),
        ("Embedding", "Linear(16, 64), compartido", "81 × 64"),
        ("Encoder × 2", "self-attention (1 cabeza) + FF 64→128→64, residual + LayerNorm", "81 × 64"),
        ("Contexto", "[h_global ; promedio de filas disponibles]", "128"),
        ("Query", "q = W_c · contexto", "64"),
        ("Puntero", "uᵣ = (W_q q)·(W_k hᵣ)/√64, máscara col. 4", "81 logits"),
        ("Salida", "softmax(u)", "81 probabilidades")]
table(s, 0.5, 1.4, 7.4, rows, [1.45, 4.35, 1.6], row_h=0.5, size=11, bold_first_col=True,
      aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER])
f1 = formula_png(r"$A=\mathrm{softmax}\!\left(\frac{QK^\top}{\sqrt{d}}\right),\quad H \leftarrow \mathrm{LN}(H + AV)$",
                 "att", fontsize=24)
f2 = formula_png(r"$q = W_c\,[\,h_{\mathrm{global}}\,;\,\bar h_{\mathrm{disp}}\,]$", "q", fontsize=24)
f3 = formula_png(r"$u_r=\frac{(W_q q)\cdot(W_k h_r)}{\sqrt{d}},\quad P(r)=\mathrm{softmax}(u)_r$", "ptr", fontsize=24)
box(s, 8.15, 1.4, 4.7, 2.9, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 8.3, 1.47, 4.4, 0.35, "Ecuaciones", size=13, bold=True, color=NAVY)
text(s, 8.3, 1.85, 4.4, 0.3, "Encoder (por capa)", size=10, color=MUTED)
image(s, f1, 8.3, 2.1, 4.4, 0.55, align="left")
text(s, 8.3, 2.7, 4.4, 0.3, "Decoder: query de contexto", size=10, color=MUTED)
image(s, f2, 8.3, 2.95, 4.4, 0.45, align="left")
text(s, 8.3, 3.42, 4.4, 0.3, "Decoder: puntero", size=10, color=MUTED)
image(s, f3, 8.3, 3.67, 4.4, 0.55, align="left")
card(s, 0.5, 5.6, 7.4, 1.3, "POR QUÉ SE RE-CODIFICA EN CADA PASO",
     [[("Las features incluyen el estado (armas libres, qⱼ, ganancia, anticipación MMR), así que H cambia en cada decisión. "
        "Costo: W pasadas del encoder por instancia (≈", {}), (f"{num(TM[MOD], 0)} ms", {"bold": True}),
       (" por instancia 10 × 8).", {})]], accent=BLUE, body_size=11)
card(s, 8.15, 4.45, 4.7, 2.45, "VS. GLIMPSE (KOOL ET AL., 2019)",
     [[("Glimpse: ", {"bold": True}), ("antes del puntero, el contexto atiende con MHA (8 cabezas) sobre H y se refina; los "
                                      "logits se recortan con C·tanh(u).", {"space_after": 4})],
      [("Nuestro decoder: ", {"bold": True}), ("contexto → puntero directo, 1 cabeza, sin recorte.", {"space_after": 4})],
      [("Encoder estático vs dinámico: ", {"bold": True}), ("Kool codifica una vez por instancia; nosotros, en cada paso.",
                                                           {})]], accent=ORANGE, body_size=10.5)

# ================================================================ 10. Configuración experimental
s = new_slide(5, "5.1 Configuración experimental", "Todos los hiperparámetros de la arquitectura y del entrenamiento",
              notes="(≈45 s) La tabla resume todos los hiperparámetros. Mantuvimos un modelo moderado: 2 capas de encoder, "
                    "1 de decoder, 1 cabeza y d = 64. El Dropout se eligió con un barrido 0 / 0,1 / 0,2 por el gap en "
                    "validación: ganó sin Dropout. Split por instancias: 6.000 para entrenar y 300 para validar; la prueba "
                    "es aparte. Adam con tasa 1e-3, batch 256, early stopping por pérdida de validación con paciencia 8.")
comma = lambda v: re.sub(r"(\d)\.(\d)", r"\1,\2", v)
arq = [(k, comma(v)) for k, v in HP["Arquitectura"].items()]
ent = [(k, comma(v)) for k, v in HP["Entrenamiento"].items()]
table(s, 0.5, 1.4, 5.3, [("Arquitectura", "Valor")] + arq, [2.55, 2.75], row_h=0.42, size=11, bold_first_col=True)
table(s, 6.0, 1.4, 6.85, [("Entrenamiento", "Valor")] + ent, [2.35, 4.5], row_h=0.48, size=10.5, bold_first_col=True)
sweep = R["dropout_sweep_t5"]
y0 = 1.4 + 0.42 * (len(arq) + 1) + 0.25
box(s, 0.5, y0, 5.3, 6.85 - y0, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 0.65, y0 + 0.06, 5.0, 0.35, "Selección del Dropout (gap rollout en validación)", size=11.5, bold=True, color=NAVY)
text(s, 0.65, y0 + 0.42, 5.0, 6.85 - y0 - 0.5,
     [[(f"Dropout {num(float(k), 1)}: ", {"bold": True}),
       (f"acc {num(100 * v['acc_val'], 1)} % · gap {num(v['gap_val'])} %", {})] for k, v in sweep.items()],
     size=11)

# ================================================================ 11. Convergencia
s = new_slide(5, "5.2 Convergencia del entrenamiento", "Pérdida y accuracy de entrenamiento y validación por época",
              notes="(≈45 s) La pérdida de entrenamiento y la de validación bajan juntas y se mantienen pegadas hasta el final: "
                    "en la mejor época la diferencia es mínima. Si hubiera sobreajuste veríamos la validación subir mientras "
                    "el entrenamiento sigue bajando. Esto es razonable porque tenemos 60.000 ejemplos y solo ~76 mil "
                    "parámetros compartidos entre todas las filas. El early stopping restaura la mejor época.")
image(s, os.path.join(FIG, "curvas.png"), 0.4, 1.35, 12.5, 4.1)
dif = CV["val_loss_best"] - CV["train_loss_best"]
k4 = [(f"época {int(CV['best_epoch'])}", "mejor época", NAVY, f"de {int(CV['epochs'])} (early stopping)"),
      (num(CV["val_loss_best"], 3), "pérdida de validación", BLUE, f"entrenamiento: {num(CV['train_loss_best'], 3)}"),
      (f"{dif:+.3f}".replace(".", ","), "diferencia val − train", GREEN if dif < 0.05 else ORANGE, "≈ 0 ⇒ sin sobreajuste"),
      (pct(100 * CV["val_acc_best"], 1), "accuracy de validación", PURPLE,
       f"entrenamiento: {pct(100 * CV['train_acc_best'], 1)}")]
for k, (v, l, col, sub) in enumerate(k4):
    kpi(s, 0.5 + k * 3.13, 5.6, 2.95, 1.25, v, l, col, sub, vsize=22)

# ================================================================ 12. Evaluación
s = new_slide(5, "5.3 Evaluación post-entrenamiento",
              "50 instancias nuevas 10 × 8 · gap = 100 · (costo / costo del experto − 1) · el ILS coincide con el óptimo verificado",
              notes="(≈1 min 15 s) Evaluamos en 50 instancias nuevas de 10 × 8, las mismas de la Tarea 4. La métrica es el gap "
                    "respecto del experto, que en este tamaño coincide con el óptimo. Comparamos con MMR, la heurística "
                    "clásica, y con un constructivo aleatorio. El modelo reduce el gap de MMR a menos de la mitad, gana en "
                    "la mayoría de las instancias y supera a la MLP de la Tarea 4. El aleatorio confirma que lo aprendido no "
                    "es trivial.")
order = [MOD, "MLP", MMR, RND, EXP]
rows = [("Método", "Costo (media ± desv.)", "Gap % (media ± desv.)", "Gap mediana", "% óptimo", "Tiempo / instancia")]
fills = []
for n in order:
    if n == "MLP":
        rows.append(("MLP Tarea 4 (ref.)", "—", f"{num(R['gap_mlp_t4'])} (solo media)", "—", "—", "—"))
        fills.append(None)
        continue
    t = T_[n]
    rows.append((n.replace(" (heurística clásica)", " (clásica)"), pm(t["mean_cost"], t["std_cost"]),
                 pm(t["mean_gap"], t["std_gap"]), pct(t["median_gap"], 2), pct(t["pct_optimal"], 0),
                 f"{num(TM[n], 1)} ms" if n in TM else "≈ 0 ms"))
    fills.append(PALE_BLUE if n == MOD else None)
table(s, 0.5, 1.4, SW - 1.0, rows, [2.5, 2.2, 2.3, 1.55, 1.4, 2.38], row_h=0.42, size=11.5, fills=fills,
      bold_first_col=True, aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 5)
image(s, os.path.join(FIG, "comparacion.png"), 0.4, 4.05, 9.3, 2.85, align="left")
w = R["model_vs_mmr"]
card(s, 9.85, 4.1, 3.0, 2.8, "LECTURA",
     [[("Modelo vs MMR: ", {"bold": True}), (f"mejor en {pct(100 * w['wins'], 0)}, igual en {pct(100 * w['ties'], 0)} de "
                                            "las instancias.", {"space_after": 4})],
      [("Gap: ", {"bold": True}), (f"{num(T_[MOD]['mean_gap'])} % vs {num(T_[MMR]['mean_gap'])} % (MMR) y "
                                  f"{num(R['gap_mlp_t4'])} % (MLP T4).", {"space_after": 4})],
      [("Aleatorio: ", {"bold": True}), (f"{num(T_[RND]['mean_gap'], 0)} % ⇒ el modelo sí aprendió.", {})]],
     accent=RED, body_size=11)

# ================================================================ 13. Desempeño por tamaño
s = new_slide(5, "5.4 Desempeño por tamaño de instancia",
              "Entrenado solo con 10 × 8 (81 filas) y evaluado sin cambios desde 5 × 4 (21 filas) hasta 30 × 20 (601 filas)",
              notes="(≈1 min) El mismo modelo, entrenado solo con 10 × 8, se aplica sin cambios a cinco tamaños, de 21 a 601 "
                    "filas. Supera a MMR en todos, y el aleatorio queda muy lejos. En 5 × 4, donde verificamos el óptimo por "
                    "fuerza bruta, es incluso mejor que en el tamaño de entrenamiento. Al crecer la instancia el gap sube y la "
                    "ventaja sobre MMR se achica: en 30 × 20 es de menos de un punto. Eso sugiere entrenar con tamaños mixtos.")
image(s, os.path.join(FIG, "tamanos.png"), 0.35, 1.35, 8.9, 3.6, align="left")
rows = [("W × T", "Modelo", "MMR", "Aleatorio")]
for l in size_labels:
    d = SZ[l]
    rows.append((l, pm(d[MOD]["mean"], d[MOD]["std"], 1), pm(d[MMR]["mean"], d[MMR]["std"], 1),
                 pm(d[RND]["mean"], d[RND]["std"], 0)))
table(s, 9.4, 1.4, 3.45, rows, [0.75, 0.95, 0.95, 0.8], row_h=0.5, size=9.5, bold_first_col=True,
      aligns=[PP_ALIGN.CENTER] * 4, header_size=10)
text(s, 9.4, 4.45, 3.45, 0.5, "gap % vs experto (media ± desv.)", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
gaps_m = [SZ[l][MOD]["mean"] for l in size_labels]
gaps_r = [SZ[l][MMR]["mean"] for l in size_labels]
beats = sum(m < r for m, r in zip(gaps_m, gaps_r))
red_ = [r - m for m, r in zip(gaps_m, gaps_r)]
i_tr = size_labels.index("10×8")
worst = size_labels[max(range(len(gaps_m)), key=lambda k: gaps_m[k])]
best_sz = size_labels[min(range(len(gaps_m)), key=lambda k: gaps_m[k])]
if best_sz == "10×8":
    trend = (f"Su mejor resultado es en el tamaño de entrenamiento (10 × 8: {num(gaps_m[i_tr])} %); fuera de él el gap "
             f"sube (peor caso {worst}: {num(max(gaps_m))} %)")
else:
    trend = (f"En 10 × 8 (entrenamiento) logra {num(gaps_m[i_tr])} %; su mejor tamaño es {best_sz} "
             f"({num(min(gaps_m))} %) y el peor {worst} ({num(max(gaps_m))} %)")
i_max = max(range(len(red_)), key=lambda k: red_[k])
adv = (f". La ventaja sobre MMR baja de {num(red_[i_max], 1)} puntos ({size_labels[i_max]}) a {num(red_[-1], 1)} en "
       f"{size_labels[-1]}." if red_[-1] < red_[i_max] else ".")
card(s, 0.5, 5.15, 6.1, 1.75, "¿SE DEGRADA AL ESCALAR?",
     [[(f"El modelo supera a MMR en {beats} de {len(size_labels)} tamaños. ", {"bold": True}), (trend + adv, {})]],
     accent=RED, body_size=11.5)
ex5 = R.get("exact_check_5x4", {}).get("5×4")
card(s, 6.75, 5.15, 6.1, 1.75, "CÓMO LEERLO",
     [[("Referencia: ", {"bold": True}), ("en 5 × 4 el ILS es óptimo" + (f" ({ex5}/30, fuerza bruta)" if ex5 else "") +
                                         "; en tamaños grandes es la mejor solución conocida.", {"space_after": 4})],
      [("Por qué generaliza: ", {"bold": True}), ("features normalizadas por instancia y pesos compartidos entre filas; "
                                                 "solo cambia la longitud de la secuencia.", {})]], accent=GRAY, body_size=11.5)

# ================================================================ 14. Discusión y conclusiones
s = new_slide(5, "5.5 Discusión y conclusiones", "Qué funcionó, qué no resultó como esperábamos y qué mejoraríamos",
              notes="(≈1 min 15 s) Funcionó: la representación por pares con features tomadas de las heurísticas, la "
                    "máscara natural y la generalización a otros tamaños; el modelo reduce el gap de MMR a menos de la "
                    "mitad. No resultó como esperábamos: el Dropout no ayudó, las primeras decisiones siguen siendo las más "
                    "difíciles y en instancias grandes la ventaja sobre MMR se achica. Mejoraríamos: agregar el glimpse "
                    "multi-cabeza al decoder, entrenar con tamaños mixtos, DAgger o RL para el distribution shift, y "
                    "reducir el costo cuadrático codificando armas y objetivos por separado.")
acc_s = R["acc_step"]
good = [[("Representación por pares: ", {"bold": True}), ("elige arma y objetivo (sin orden fijo); máscara natural.",
                                                          {"space_after": 4})],
        [("Features de las heurísticas: ", {"bold": True}),
         (f"accuracy {pct(100 * R['val_acc_final'], 1)} vs {pct(100 * sum(R['acc_mmr_step']) / len(R['acc_mmr_step']), 1)} "
          "de la regla MMR.", {"space_after": 4})],
        [("Calidad: ", {"bold": True}), (f"gap {num(T_[MOD]['mean_gap'])} % vs {num(T_[MMR]['mean_gap'])} % de MMR, "
                                         f"óptimo en {pct(T_[MOD]['pct_optimal'], 0)} de las instancias.",
                                         {"space_after": 4})],
        [("Generaliza ", {"bold": True}), ("de 21 a 601 filas sin reentrenar.", {})]]
bad = [[("Dropout no ayudó: ", {"bold": True}), ("con 60 mil ejemplos y pesos compartidos no hubo sobreajuste.",
                                                 {"space_after": 4})],
       [("Primeras decisiones: ", {"bold": True}), (f"accuracy {pct(100 * acc_s[0], 0)} en el paso 0 (80 acciones) vs "
                                                    f"{pct(100 * acc_s[-1], 0)} en el último.", {"space_after": 4})],
       [(f"En {size_labels[-1]} ", {"bold": True}), (f"la ventaja sobre MMR cae a {num(red_[-1], 1)} puntos "
                                                     f"(gap {num(gaps_m[-1])} % vs {num(gaps_r[-1])} %).", {"space_after": 4})],
       [("Costo: ", {"bold": True}), (f"{num(TM[MOD], 0)} ms por instancia: re-codificar en cada paso y atención O(n²).", {})]]
better = [[("Decoder con Glimpse ", {"bold": True}), ("y atención multi-cabeza (Kool et al.).", {"space_after": 4})],
          [("Entrenar con tamaños mixtos ", {"bold": True}), ("(5 × 4 a 20 × 15) para generalizar mejor.", {"space_after": 4})],
          [("DAgger o RL ", {"bold": True}), ("para corregir el distribution shift tras un error temprano.", {"space_after": 4})],
          [("Encoder de W + T elementos ", {"bold": True}), ("con puntero bilineal, para bajar el costo cuadrático.", {})]]
for k, (t, body, col, fill) in enumerate([("¿QUÉ FUNCIONÓ BIEN?", good, GREEN, PALE_GREEN),
                                          ("¿QUÉ NO RESULTÓ?", bad, ORANGE, PALE_ORANGE),
                                          ("¿QUÉ MEJORARÍAMOS?", better, BLUE, PALE_BLUE)]):
    card(s, 0.5 + k * (cw + 0.2), 1.4, cw, 3.55, t, body, accent=col, title_size=14, body_size=12.5, fill=fill)
box(s, 0.5, 5.15, SW - 1.0, 1.7, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 0.75, 5.2, SW - 1.5, 1.6,
     [[("Conclusión: ", {"bold": True, "color": RGBColor(0xFF, 0xC0, 0x7A)}),
       (f"un encoder–decoder con atención, entrenado solo por imitación, aprende una política constructiva que reduce el gap de "
        f"MMR de {num(T_[MMR]['mean_gap'])} % a {num(T_[MOD]['mean_gap'])} % y que se usa sin cambios en instancias de otros "
        "tamaños. La siguiente etapa natural es el glimpse y el entrenamiento por refuerzo.", {})]],
     size=15, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

assert len(prs.slides) == TOTAL, len(prs.slides)
for sl, notes in SLIDES:
    sl.notes_slide.notes_text_frame.text = notes
prs.save(OUT)
for f in os.listdir(TMP):
    os.remove(os.path.join(TMP, f))
os.rmdir(TMP)
print("Presentación guardada en", OUT)
