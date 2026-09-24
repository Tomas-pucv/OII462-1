"""Genera la presentación de la Tarea 4 (WTAP) a partir de las figuras y resultados del notebook.

Uso (desde la raíz del repositorio):
    python scripts/make_presentation.py
    soffice --headless --convert-to pdf Tarea4_WTAP_presentacion.pptx
"""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figuras")
OUT = os.path.join(ROOT, "Tarea4_WTAP_presentacion.pptx")
R = json.load(open(os.path.join(FIG, "resultados.json"), encoding="utf-8"))

FONT = "Arial"
NAVY = RGBColor(0x1B, 0x2A, 0x41)
INK = RGBColor(0x2E, 0x35, 0x40)
MUTED = RGBColor(0x6B, 0x75, 0x85)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
ORANGE = RGBColor(0xFF, 0x7F, 0x0E)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
PURPLE = RGBColor(0x94, 0x67, 0xBD)
RED = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, shape=None):
    """runs: str, o lista de párrafos; cada párrafo es str o lista de (texto, {opciones})."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = runs if isinstance(runs, list) else [runs]
    for k, par in enumerate(paragraphs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        pieces = par if isinstance(par, list) else [(par, {})]
        for txt, opt in pieces:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
        if isinstance(par, list) and par and "space_after" in par[0][1]:
            p.space_after = Pt(par[0][1]["space_after"])
    return shape


def image(slide, path, x, y, w=None, h=None, align="center"):
    """Inserta la imagen ajustada a la caja (x, y, w, h) manteniendo la proporción."""
    iw, ih = Image.open(path).size
    ratio = iw / ih
    if w is not None and h is not None:
        if w / h > ratio:
            nw, nh = h * ratio, h
        else:
            nw, nh = w, w / ratio
        ox = x + {"center": (w - nw) / 2, "left": 0, "right": w - nw}[align]
        oy = y + (h - nh) / 2
    elif w is not None:
        nw, nh, ox, oy = w, w / ratio, x, y
    else:
        nw, nh, ox, oy = h * ratio, h, x, y
    return slide.shapes.add_picture(path, Inches(ox), Inches(oy), Inches(nw), Inches(nh))


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def header(slide, n, title, subtitle=None):
    box(slide, 0, 0, SW, SH, fill=WHITE)
    box(slide, 0.5, 0.42, 0.09, 0.62, fill=BLUE)
    text(slide, 0.72, 0.3, 11.5, 0.6, title, size=28, bold=True, color=NAVY)
    if subtitle:
        text(slide, 0.74, 0.86, 11.8, 0.4, subtitle, size=14, color=MUTED)
    text(slide, 0.5, 7.05, 8, 0.3, "OII462 · Tarea 4 · Imitando un algoritmo de trayectoria para el WTAP",
         size=10, color=MUTED)
    text(slide, SW - 1.3, 7.05, 0.8, 0.3, f"{n} / 6", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=15, body_size=12.5, fill=LIGHT):
    box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(slide, x, y + 0.15, 0.07, h - 0.3, fill=accent)
    text(slide, x + 0.22, y + 0.1, w - 0.35, 0.4, title, size=title_size, bold=True, color=accent)
    text(slide, x + 0.22, y + 0.5, w - 0.35, h - 0.6, body, size=body_size, color=INK, spacing=1.12)


def kpi(slide, x, y, w, h, value, label, color, sub=None):
    box(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(slide, x, y, w, 0.08, fill=color)
    text(slide, x, y + 0.14, w, 0.55, value, size=26, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(slide, x + 0.05, y + 0.7, w - 0.1, 0.35, label, size=11.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
    if sub:
        text(slide, x + 0.05, y + 0.98, w - 0.1, 0.3, sub, size=10, color=MUTED, align=PP_ALIGN.CENTER)


def formula_png(latex, path, fontsize=26, color="#1B2A41"):
    fig = plt.figure(figsize=(8, 1))
    fig.text(0, 0.5, latex, fontsize=fontsize, color=color, va="center")
    fig.savefig(path, dpi=250, bbox_inches="tight", pad_inches=0.05, transparent=True)
    plt.close(fig)
    return path


def pct(v, d=1):
    return f"{v:.{d}f}%".replace(".", ",")


def num(v, d=3):
    return f"{v:.{d}f}".replace(".", ",")


T = R["test"]
M_NAME, B_NAME = "Modelo MLP", "Modelo base (sin anticipación)"
MMR, SEQ, EXP = "MMR (heurística clásica)", "Greedy secuencial", "Experto (ILS)"
tmp = os.path.join(FIG, "_tmp")
os.makedirs(tmp, exist_ok=True)

# ================================================================ 1. Problema
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=WHITE)
box(s, 0, 0, 4.7, SH, fill=NAVY)
text(s, 0.55, 0.7, 3.8, 0.4, "OII462 · TAREA 4", size=13, bold=True, color=RGBColor(0x8F, 0xB8, 0xDE))
text(s, 0.55, 1.15, 3.9, 2.4, "Imitando un algoritmo de trayectoria para el WTAP",
     size=31, bold=True, color=WHITE, spacing=1.0)
text(s, 0.55, 3.45, 3.8, 1.2, "Una red de capas densas (MLP) que aprende a asignar armas a objetivos "
     "imitando paso a paso a un experto óptimo", size=15, color=RGBColor(0xD6, 0xE2, 0xEE), spacing=1.15)
box(s, 0.55, 4.95, 0.9, 0.05, fill=ORANGE)
text(s, 0.55, 5.15, 3.8, 0.4, "[Nombre(s) del grupo]", size=16, bold=True, color=WHITE)
text(s, 0.55, 5.55, 3.8, 0.4, "Optimización combinatoria · 2026", size=12, color=RGBColor(0xB5, 0xC4, 0xD4))

text(s, 5.2, 0.42, 7.9, 0.5, "El problema: Weapon-Target Assignment", size=23, bold=True, color=NAVY)
text(s, 5.2, 1.0, 7.8, 0.75,
     [[("W = 10 armas", {"bold": True}), (" y ", {}), ("T = 8 objetivos", {"bold": True}),
       (". El objetivo j vale ", {}), ("Vⱼ", {"bold": True}), (" y el arma i lo destruye con probabilidad ", {}),
       ("pᵢⱼ", {"bold": True}), (". Cada arma se asigna a un solo objetivo.", {})]], size=14)
fpath = formula_png(r"$\min_x\; f(x)=\sum_{j=1}^{T} V_j \prod_{i=1}^{W}(1-p_{ij})^{x_{ij}}$",
                    os.path.join(tmp, "formula.png"))
box(s, 5.2, 1.85, 7.6, 0.95, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
image(s, fpath, 5.35, 1.93, 4.7, 0.8, align="left")
text(s, 10.05, 1.93, 2.7, 0.8, "valor esperado que sobrevive (se minimiza)", size=12, color=MUTED,
     anchor=MSO_ANCHOR.MIDDLE)
image(s, os.path.join(FIG, "instancia.png"), 5.2, 3.0, 7.7, 2.95)
facts = [("NP-hard", "objetivo no lineal"), ("8¹⁰ ≈ 10⁹", "asignaciones posibles"),
         ("Constructivo", "1 arma por paso, orden fijo")]
for k, (big, small) in enumerate(facts):
    x = 5.2 + k * 2.6
    box(s, x, 6.12, 2.45, 0.8, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, x, 6.14, 2.45, 0.42, big, size=17, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x, 6.53, 2.45, 0.35, small, size=11, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 2. Estado y acción
s = prs.slides.add_slide(BLANK)
header(s, 2, "Estado y acción: ¿a qué objetivo va la próxima arma?",
       "Las armas se asignan en orden fijo (0, 1, …, 9); en cada paso la red elige el objetivo de la arma actual")
card(s, 0.5, 1.45, 4.1, 1.55, "ACCIÓN  →  8 clases",
     [[("(\"constructive\", j)", {"bold": True}), (": asignar la arma actual al objetivo j. ", {}),
       ("El target de la red es el índice j.", {})]], accent=ORANGE)
card(s, 0.5, 3.15, 4.1, 1.55, "ESTADO  →  matriz 8 × 15",
     [[("Fila j = objetivo j", {"bold": True}), (" durante todo el episodio; 15 características que "
       "describen al objetivo y su situación en el paso actual.", {})]], accent=BLUE)
card(s, 0.5, 4.85, 4.1, 2.05, "¿QUÉ MIRA LA RED?",
     [[("Actuales (0–9): ", {"bold": True, "color": BLUE}),
       ("valor, prob. de sobrevivir qⱼ, efectividad de la arma actual pₜⱼ, ", {}),
       ("ganancia marginal Vⱼ·qⱼ·pₜⱼ", {"bold": True}), (", armas ya asignadas…", {})],
      [("Anticipación (10–14): ", {"bold": True, "color": PURPLE}),
       ("¿hay una arma futura mejor para j? (arrepentimiento, rango) y cómo quedaría j si se completa con MMR.",
        {})]], accent=PURPLE, body_size=12)
image(s, os.path.join(FIG, "estado.png"), 4.85, 1.45, 8.0, 2.9)
box(s, 4.95, 4.55, 4.25, 2.35, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 5.15, 4.65, 3.9, 2.2,
     [[("Ejemplo real: ", {"bold": True}), ("estado con 4 armas asignadas (arma actual = 4).", {})],
      [("El recuadro rojo marca el objetivo que elige el experto; la línea blanca separa las características "
        "actuales de las de anticipación.", {})],
      [("Todo está normalizado (valores ÷ máximo de la instancia).", {"color": MUTED, "space_after": 6})],
      [("→ Ninguna característica sola da la respuesta: la red aprende a combinarlas.",
        {"bold": True, "color": BLUE})]],
     size=12, spacing=1.12)
ej = R["ejemplo_estado"]
box(s, 9.4, 4.55, 3.45, 2.35, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 9.55, 4.62, 3.2, 0.4, "¿A qué objetivo iría la arma 4?", size=13, bold=True, color=NAVY)
for k, (who, tgt, col) in enumerate([("Greedy (mayor ganancia)", ej["greedy"], GRAY),
                                     ("Completar con MMR", ej["mmr"], ORANGE),
                                     ("Experto (lo que se imita)", ej["experto"], GREEN)]):
    yy = 5.1 + k * 0.58
    text(s, 9.55, yy, 2.45, 0.45, who, size=11.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    chip = box(s, 11.95, yy + 0.04, 0.75, 0.38, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, 0, 0, 0, 0, f"o{tgt}", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=chip)

# ================================================================ 3. Datos de imitación y modelo
s = prs.slides.add_slide(BLANK)
header(s, 3, "Datos de imitación y modelo",
       "El experto resuelve cada instancia; su solución se reproduce paso a paso para obtener pares (estado → acción)")
steps = [
    ("1", "Instancias aleatorias", f"{R['n_train_instances']:,}".replace(",", ".") +
     " de entrenamiento + 300 de validación\nVⱼ ~ U{25..100}, pᵢⱼ ~ U(0,6; 0,9)", NAVY),
    ("2", "Experto: Búsqueda Local Iterada", "inicio MMR + búsqueda local reassign/swap + perturbación.\n"
     f"Óptimo en {R['ils_optimal_hits']}/{R['ils_optimal_checked']} instancias verificadas con método exacto",
     GREEN),
    ("3", "Replay de la trayectoria", "solución vacía → asignar arma 0, 1, …, 9 como el experto "
     "con env.state_transition", BLUE),
    ("4", "Pares (estado → acción)", f"{R['n_train_samples']:,}".replace(",", ".") +
     " ejemplos de entrenamiento\nX: state2vec (8×15)   Y: objetivo elegido", ORANGE),
]
bw, gapx, y0 = 2.83, 0.33, 1.5
for k, (n, title, body, col) in enumerate(steps):
    x = 0.5 + k * (bw + gapx)
    box(s, x, y0, bw, 2.05, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    c = box(s, x + 0.18, y0 + 0.18, 0.5, 0.5, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, 0, 0, 0, 0, n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         shape=c)
    text(s, x + 0.78, y0 + 0.14, bw - 0.9, 0.6, title, size=14, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.18, y0 + 0.8, bw - 0.3, 1.2, body, size=11.5, color=INK)
    if k < 3:
        arrow(s, x + bw + 0.03, y0 + 1.02, x + bw + gapx - 0.03, y0 + 1.02, color=MUTED)

# diagrama de la MLP
cfg = R["final_config"]
layers = [("Estado", f"8 × {R['final_n_features']}", NAVY), ("Aplanar", f"{8 * R['final_n_features']}", NAVY)]
for hsize in cfg["hidden_sizes"]:
    layers.append(("Densa + ReLU", f"{hsize}" + (f"\nDropout {num(cfg['dropout'], 1)}" if cfg["dropout"] else ""),
                   BLUE))
layers += [("Logits", "8 (uno por objetivo)", ORANGE), ("Softmax", "P(arma → objetivo j)", GREEN)]
text(s, 0.5, 3.85, 6, 0.4, "Modelo: MLP de capas densas (PyTorch)", size=17, bold=True, color=NAVY)
n = len(layers)
lw = 1.45
lg = (SW - 1.0 - n * lw) / (n - 1)
for k, (name, detail, col) in enumerate(layers):
    x = 0.5 + k * (lw + lg)
    hgt = 1.55 if col == BLUE else 1.25
    yy = 4.45 + (1.55 - hgt) / 2
    b = box(s, x, yy, lw, hgt, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 0, 0, 0, 0, [[(name, {"bold": True, "size": 13})], [(detail, {"size": 11.5})]], color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
    if k < n - 1:
        arrow(s, x + lw + 0.02, 5.22, x + lw + lg - 0.02, 5.22, color=MUTED)
box(s, 0.5, 6.25, SW - 1.0, 0.65, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
text(s, 0.7, 6.28, SW - 1.4, 0.6,
     [[("Entrenamiento: ", {"bold": True}), ("entropía cruzada · Adam (lr = 0,001) · batch "
       f"{cfg['batch_size']} · ", {}), ("early stopping", {"bold": True}),
       (" (patience 10, restaura la mejor época). En el rollout, la red reemplaza a la función de evaluación "
        "del GreedyAgent de la librería del curso.", {})]], size=12, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 4. Entrenamiento
s = prs.slides.add_slide(BLANK)
header(s, 4, "Entrenamiento: loss y accuracy por época",
       "Early stopping detiene el entrenamiento cuando la pérdida de validación deja de mejorar")
curvas = os.path.join(FIG, "curvas_mejor.png") if os.path.exists(os.path.join(FIG, "curvas_mejor.png")) \
    else os.path.join(FIG, "curvas.png")
image(s, curvas, 0.4, 1.4, 12.5, 3.75)
kpi(s, 0.6, 5.4, 2.6, 1.4, pct(100 * R["val_acc_final"]), "accuracy de validación", BLUE,
    "decisiones iguales al experto")
kpi(s, 3.4, 5.4, 2.6, 1.4, f"época {R['best_epoch_final']}", "mejor época", NAVY,
    f"se detuvo en la época {R['epochs_final'] - 1}")
kpi(s, 6.2, 5.4, 2.6, 1.4, pct(100 * R["val_acc_base"]), "acc. modelo base", GRAY, "sin anticipación")
box(s, 9.0, 5.4, 3.85, 1.4, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 9.15, 5.45, 3.6, 1.3,
     [[("Lectura: ", {"bold": True}), ("la pérdida de entrenamiento sigue bajando mientras la de validación "
       "se estanca → comienza el sobreajuste y se restauran los pesos de la mejor época.", {})]],
     size=11.5, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 5. Experimentos
s = prs.slides.add_slide(BLANK)
header(s, 5, "Experimentos: ¿qué necesita saber la red?",
       "Ablación de características e hiperparámetros, evaluando accuracy Y calidad del rollout (100 instancias de validación)")
image(s, os.path.join(FIG, "experimentos.png"), 0.4, 1.35, 12.5, 3.9)
E = R["experiments"]
names = list(E)
rows = [("Config.", "Características", "Red", "Acc. val", "Gap rollout")]
FEAT_LABEL = {"10": "10 (sin anticipación)", "12": "12 (+ comparativas)", "15": "15 (completas)"}
for nme in names:
    code_, rest = nme.split(" · ", 1)
    feats, _, net = rest.partition(",")
    rows.append((code_, FEAT_LABEL[feats.split()[0]], (net.strip() or "[512,256,128], dp 0.2").replace("0.", "0,"),
                 pct(100 * E[nme]["val_acc"]), pct(E[nme]["gap_val"])))
tbl = s.shapes.add_table(len(rows), 5, Inches(0.5), Inches(5.35), Inches(7.6), Inches(1.6)).table
widths = [0.8, 2.6, 2.2, 1.0, 1.0]
for c, w in enumerate(widths):
    tbl.columns[c].width = Inches(w)
for r_, row in enumerate(rows):
    tbl.rows[r_].height = Inches(1.6 / len(rows))
    for c, val in enumerate(row):
        cell = tbl.cell(r_, c)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        cell.fill.solid()
        is_best = r_ > 0 and names[r_ - 1] == R["final_model"]
        cell.fill.fore_color.rgb = NAVY if r_ == 0 else (RGBColor(0xDC, 0xEA, 0xF7) if is_best else
                                                          (LIGHT if r_ % 2 else WHITE))
        text(s, 0, 0, 0, 0, val, size=10, bold=(r_ == 0 or is_best), color=WHITE if r_ == 0 else INK,
             align=PP_ALIGN.LEFT if c in (1, 2) else PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=cell)
best_gap = E[R["final_model"]]["gap_val"]
base_gap = E[names[0]]["gap_val"]
card(s, 8.35, 5.35, 4.5, 1.6, "CONCLUSIÓN",
     [[("Las características de anticipación son lo que más importa: el gap baja de ", {}),
       (pct(base_gap), {"bold": True}), (" a ", {}), (pct(best_gap), {"bold": True}),
       (f" (MMR: {pct(R['gap_mmr_val'])}). Cambiar la red o el Dropout influye menos, y con igual "
        "accuracy el gap puede variar: hay que evaluar el rollout.", {})]],
     accent=PURPLE, body_size=11.5)

# ================================================================ 6. Resultados
s = prs.slides.add_slide(BLANK)
header(s, 6, f"Resultados en {R['n_test_instances']} instancias nuevas",
       "Gap promedio respecto al experto (menor es mejor) · mismo agente greedy; solo cambia quién evalúa las acciones")
kpis = [(M_NAME, "Modelo MLP", BLUE), (MMR, "MMR (clásica)", ORANGE), (SEQ, "Greedy secuencial", GRAY),
        (EXP, "Experto (ILS)", GREEN)]
for k, (key, label, col) in enumerate(kpis):
    kpi(s, 0.5 + k * 3.13, 1.4, 2.95, 1.3, pct(T[key]["mean_gap"]), label, col,
        f"costo medio {num(T[key]['mean_cost'], 1)} · óptimo en {pct(T[key]['pct_optimal'], 0)}")
image(s, os.path.join(FIG, "comparacion.png"), 0.4, 2.85, 12.5, 3.3)
wins = R["model_vs_mmr"]
box(s, 0.5, 6.2, SW - 1.0, 0.72, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
text(s, 0.7, 6.22, SW - 1.4, 0.68,
     [[("Conclusión: ", {"bold": True}),
       (f"La MLP gana a MMR en {pct(100 * wins['wins'], 0)} de las instancias (empata en "
        f"{pct(100 * wins['ties'], 0)}) y reduce el gap del greedy secuencial de {pct(T[SEQ]['mean_gap'])} a "
        f"{pct(T[M_NAME]['mean_gap'])}, respetando el mismo orden de armas. Sin las características de anticipación el modelo queda en "
        f"{pct(T[B_NAME]['mean_gap'])}: lo que la red ve del estado es clave.", {})]],
     size=12, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
for f in os.listdir(tmp):
    os.remove(os.path.join(tmp, f))
os.rmdir(tmp)
print("Presentación guardada en", OUT)
