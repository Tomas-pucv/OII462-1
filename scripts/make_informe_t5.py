"""Genera el informe de la Tarea 5 (puntos 1-3) a partir de las figuras y resultados del notebook.

Uso (desde la raíz del repositorio):
    python scripts/make_informe_t5.py
    soffice --headless --convert-to pdf Tarea5_WTAP_informe.docx
"""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figuras_t5")
OUT = os.path.join(ROOT, "Tarea5_WTAP_informe.docx")
R = json.load(open(os.path.join(FIG, "resultados.json"), encoding="utf-8"))
TMP = os.path.join(FIG, "_tmp_informe")
os.makedirs(TMP, exist_ok=True)

FONT = "Arial"
NAVY = RGBColor(0x1B, 0x2A, 0x41)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
MUTED = RGBColor(0x6B, 0x75, 0x85)
INK = RGBColor(0x2E, 0x35, 0x40)
ROLE_FILL = {"propia": "DCEAF7", "flag": "F9DEDC", "relativa": "DFF0DF", "global": "EEEEEE"}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)      # carta
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
TEXT_W = 21.59 - 4.4

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.12
for name, size in [("Heading 1", 15), ("Heading 2", 12)]:
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
    h.font.color.rgb = NAVY
    rfonts = h.element.rPr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); h.element.rPr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FONT)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    h.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- helpers
def para(runs, style=None, align=None, size=None, space_after=None, keep=False):
    """runs: str o lista de (texto, {bold, italic, color, font, size})."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep:
        p.paragraph_format.keep_with_next = True
    for txt, opt in ([(runs, {})] if isinstance(runs, str) else runs):
        r = p.add_run(txt)
        r.bold = opt.get("bold", False)
        r.italic = opt.get("italic", False)
        if "color" in opt:
            r.font.color.rgb = opt["color"]
        if "font" in opt:
            r.font.name = opt["font"]
        if size or "size" in opt:
            r.font.size = Pt(opt.get("size", size))
    return p


def bullet(runs):
    p = para(runs, style="List Bullet", space_after=2)
    p.paragraph_format.left_indent = Cm(0.9)
    return p


def c(txt):
    return (txt, {"font": "Liberation Mono", "size": 9.5})


def b(txt):
    return (txt, {"bold": True})


def i(txt):
    return (txt, {"italic": True})


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_text(cell, runs, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for txt, opt in ([(runs, {})] if isinstance(runs, str) else runs):
        r = p.add_run(txt)
        r.font.size = Pt(opt.get("size", size))
        r.bold = opt.get("bold", bold)
        r.italic = opt.get("italic", False)
        if "font" in opt:
            r.font.name = opt["font"]
        if color is not None:
            r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(rows, widths, header_fill="1B2A41", fills=None, size=9, aligns=None):
    t = doc.add_table(rows=len(rows), cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c_, w in enumerate(widths):
        t.columns[c_].width = Cm(w)
    for r_, row in enumerate(rows):
        for c_, val in enumerate(row):
            cell = t.cell(r_, c_)
            cell.width = Cm(widths[c_])
            al = (aligns[c_] if aligns else WD_ALIGN_PARAGRAPH.LEFT)
            if r_ == 0:
                cell_text(cell, val, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=al)
                shade(cell, header_fill)
            else:
                cell_text(cell, val, size=size, align=al)
                if fills and fills[r_ - 1]:
                    shade(cell, fills[r_ - 1])
    for row in t.rows:                                   # evita que una fila se parta entre páginas
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); trPr.append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    para([(caption, {"italic": True, "color": MUTED})], align=WD_ALIGN_PARAGRAPH.CENTER, size=9, space_after=8)


def formula(latex, name, height_cm=1.0, fontsize=17):
    path = os.path.join(TMP, f"{name}.png")
    fig = plt.figure(figsize=(8, 0.6))
    fig.text(0, 0.5, latex, fontsize=fontsize, color="#1B2A41", va="center")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(path, height=Cm(height_cm))


def code_block(lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(TEXT_W - 0.4)
    cell = t.cell(0, 0)
    cell.width = Cm(TEXT_W - 0.4)
    shade(cell, "F3F5F8")
    cell.text = ""
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        r.font.name = "Liberation Mono"
        r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pct(v, d=1):
    return f"{v:.{d}f}".replace(".", ",") + " %"


def miles(n):
    return f"{n:,}".replace(",", ".")


D = R["demo"]
T_ = R["test"]
W, T = 10, 8
N_ROWS = R["n_rows"]

# ================================================================ portada / encabezado
para([("OII462 · Tarea 5", {"bold": True, "color": BLUE, "size": 10})], space_after=0)
para([("Datos de imitación y modelo de atención con puntero para el WTAP", {"bold": True, "color": NAVY, "size": 19})],
     space_after=2)
para([("Informe: descomposición de soluciones, codificación de la entrada y de la salida (puntos 1–3)",
       {"color": MUTED, "size": 11})], space_after=2)
para([("[Nombre(s) del grupo]", {"bold": True}), ("  ·  Optimización combinatoria · 2026", {"color": MUTED})],
     space_after=10)

para([b("Problema. "),
      ("En el Weapon-Target Assignment Problem (WTAP) hay W armas y T objetivos; el objetivo j vale Vⱼ y el arma i "
       "lo destruye con probabilidad pᵢⱼ. Cada arma se asigna a lo más a un objetivo (un objetivo puede recibir "
       "varias) y se minimiza el valor esperado que sobrevive:", {})])
formula(r"$\min_x\; Z=\sum_{j=1}^{T} V_j \prod_{i=1}^{W}(1-p_{ij})^{x_{ij}}, \qquad q_j=\prod_{i}(1-p_{ij})^{x_{ij}}$",
        "objetivo", height_cm=1.05)
para([("Llamamos qⱼ a la probabilidad de que el objetivo j sobreviva con las armas asignadas hasta ahora. "
       "Las instancias siguen la literatura (Vⱼ ~ U{25,…,100}, pᵢⱼ ~ U(0,6; 0,9)); se entrena con ", {}),
      b("W = 10 y T = 8"), (", pero el modelo acepta cualquier tamaño.", {})])
para([b("Qué cambia respecto de la Tarea 4. "),
      ("La MLP de la Tarea 4 asignaba las armas en orden fijo (arma 0, 1, …) y elegía el objetivo de la arma actual. "
       "Tenía pesos distintos por posición, tamaño fijo y no podía decidir qué arma usar primero. Con un encoder de "
       "self-attention y un puntero, cada ", {}), b("acción posible es un par (arma i → objetivo j)"),
      (" representado por una fila de la matriz de entrada, y la decisión puntual del modelo es: ", {}),
      b("dado el estado actual, ¿qué arma asigno ahora y a qué objetivo?"),
      (" Es la misma pregunta que se hace MMR (Maximum Marginal Return), la heurística constructiva clásica del WTAP.",
       {})])

# ================================================================ 1. Descomposición
doc.add_heading("1. Descomposición de las soluciones en decisiones", level=1)
doc.add_heading("Algoritmo tradicional (experto) y tipo de acciones", level=2)
para([("El experto es la ", {}), b("Búsqueda Local Iterada (ILS)"),
      (" de la Tarea 4: parte de la solución de MMR, aplica búsqueda local best-improvement con los vecindarios "
       "reassign y swap (deltas de costo vectorizados), perturba 3 armas al azar y acepta si mejora (o con "
       "probabilidad 5 %). En instancias 10×8 alcanza el óptimo en todas las instancias verificadas con un método "
       "exacto meet-in-the-middle.", {})])
para([("Usamos ", {}), b("acciones constructivas"),
      (": cada acción agrega un par (arma libre i → objetivo j) a la solución parcial. No usamos la secuencia de "
       "movimientos de búsqueda local del experto porque esa trayectoria recorre cientos de movimientos ligados a "
       "perturbaciones aleatorias y no es una política reproducible a partir del estado. Una construcción, en "
       "cambio, tiene exactamente W decisiones y cada una depende solo del estado parcial, igual que en el ejemplo del "
       "TSP. Además, representar la acción como un par (y no como \"el objetivo de la arma actual\") elimina la "
       "restricción de orden fijo: el modelo puede elegir qué arma usar, como MMR.", {})])

doc.add_heading("Regla de orden: de un conjunto de pares a una secuencia", level=2)
para([("La solución del experto es un ", {}), b("conjunto"), (" de W pares (i, a*ᵢ) sin orden. Para descomponerla "
      "necesitamos decidir en qué orden se agregan los pares. Usamos el criterio de MMR: en cada paso, entre las armas "
      "aún libres, se toma el par del experto con ", {}), b("mayor ganancia marginal"), (" en el estado actual:", {})])
formula(r"$i^\star=\arg\max_{i\ \mathrm{libre}}\; V_{a^*_i}\; q_{a^*_i}\; p_{i\,a^*_i}, \qquad"
        r"\mathrm{acci\'on\ escogida} = (i^\star,\ a^*_{i^\star})$", "orden", height_cm=1.0)
para([("La regla es determinista (una sola etiqueta por estado) y cualquier orden reconstruye exactamente la solución "
       "del experto. Como coincide con el criterio de MMR, las decisiones con grandes ganancias van primero y la política "
       "aprendida puede leerse como “MMR corregido por el experto”: en la mayoría de los pasos el par escogido es el mismo "
       "que elegiría MMR, y el modelo debe aprender a detectar cuándo no lo es.", {})])

doc.add_heading("Tuplas (estado, acciones posibles, acción escogida)", level=2)
bullet([b("Estado inicial: "), ("asignación vacía, todas las armas libres (", {}), c("assignment = [-1]*W"),
        ("). Es el análogo de ", {}), c("visited=[0]"), (" en el TSP.", {})])
bullet([b("Acciones posibles: "), ("todos los pares (arma libre, objetivo), generados por ", {}),
        c("env.gen_actions(state, \"constructive\")"), (". Con k armas libres son k·T (80 en el primer paso, 8 en el "
        "último).", {})])
bullet([b("Acción escogida: "), ("el par del experto que indica la regla de orden. Luego se aplica ", {}),
        c("env.state_transition(state, (\"constructive\", (i, j)))"), (", que fija xᵢⱼ = 1 y actualiza qⱼ ← qⱼ(1 − pᵢⱼ).",
                                                                        {})])
bullet([b("Término: "), ("cuando las W armas están asignadas, así que hay W tuplas por instancia. A diferencia del TSP no "
        "se corta antes: incluso la última arma tiene T destinos posibles y esa decisión no es trivial.", {})])
code_block([
    "def decompose(inst, expert_assignment):",
    "    a = expert_assignment",
    "    state = WTAP_State(inst)                          # estado inicial: nada asignado",
    "    while not state.is_complete:                      # término: todas las armas asignadas",
    "        free = state.free_weapons",
    "        gains = [V[a[i]] * state.q[a[i]] * P[i, a[i]] for i in free]",
    "        i = free[argmax(gains)]                       # regla de orden (criterio MMR)",
    "        chosen = (\"constructive\", (i, a[i]))",
    "        yield deepcopy(state), list(env.gen_actions(state, \"constructive\")), chosen",
    "        env.state_transition(state, chosen)",
])

doc.add_heading("Ejemplo concreto", level=2)
para([("Instancia de ejemplo (10 armas × 8 objetivos). El experto encuentra la asignación ", {}),
      c(str(D["expert_solution"])), (f" (arma → objetivo) con costo {D['expert_cost']:.2f}".replace(".", ",", 1)
                                     + ". Su descomposición en 10 decisiones es:", {})], keep=True)
rows = [("Paso", "Armas libres", "Acciones posibles", "Acción escogida (experto)", "MMR elegiría")]
fills = []
for s in D["steps"]:
    same = s["mmr"] == s["chosen"]
    rows.append((str(s["k"]), str(s["free"]), str(s["n_actions"]), f"arma {s['chosen'][0]} → objetivo {s['chosen'][1]}",
                 f"arma {s['mmr'][0]} → objetivo {s['mmr'][1]}" + ("" if same else "  (≠)")))
    fills.append(None if same else "FFF2CC")
C = WD_ALIGN_PARAGRAPH.CENTER
table(rows, [1.4, 2.4, 3.0, 5.0, 5.0], fills=fills, aligns=[C, C, C, C, C])
n_diff = sum(1 for s in D["steps"] if s["mmr"] != s["chosen"])
para([(f"En {n_diff} de los 10 pasos (en amarillo) el par del experto no es el de mayor ganancia: son justamente "
       "las decisiones que MMR toma mal y que el modelo debe aprender a corregir. En todo el conjunto de validación, "
       f"MMR coincide con el experto en {pct(100 * sum(R['acc_mmr_step']) / len(R['acc_mmr_step']))} de los pasos.",
       {})])
figure(os.path.join(FIG, "descomposicion.png"), TEXT_W,
       "Figura 1. Descomposición de la solución del experto: en cada paso, las líneas azules son los pares ya asignados, "
       "la roja es la acción escogida y los círculos vacíos son las armas libres. Tamaño del objetivo ∝ valor; "
       "color = probabilidad de sobrevivir.")

# ================================================================ 2. state2vec
doc.add_heading("2. Codificación de la entrada (state2vec)", level=1)
para([("La entrada no se aplana: es una matriz de forma ", {}), b(f"(1 + W·T, 16)"),
      (f" ({N_ROWS} × 16 en las instancias de entrenamiento), con dtype float32. Cada fila se procesa con los mismos "
       "pesos:", {})])
bullet([b("Fila 0: fila global"), (" (tipo [CLS]). Resume el estado completo y forma la query de decisión. Se identifica por "
        "el flag ", {}), c("es_global"), (" (columna 3) y además está siempre en la posición 0, como la ciudad inicial del "
        "TSP. Nunca es una acción válida.", {})])
bullet([b("Fila 1 + i·T + j: par (arma i → objetivo j)."),
        (" Describe a la vez el estado y la acción “asignar el arma i al objetivo j”. La correspondencia fila ↔ par es fija "
         "durante todo el episodio.", {})])
para([("Las columnas replican los tres roles del notebook del TSP. Las características relativas al estado se tomaron de "
       "lo que miran las heurísticas constructivas del WTAP: la ", {}), b("ganancia marginal"),
      (" de MMR, la comparación con otras armas libres (arrepentimiento o regret) y la anticipación que se obtiene al "
       "completar la solución con MMR (que en la Tarea 4 resultó la información más valiosa). Sea F el conjunto de "
       "armas libres y gᵢⱼ = Vⱼ qⱼ pᵢⱼ.", {})])

FEAT = [
    ("0", "p_ij", "propia", "pᵢⱼ", "efectividad del arma contra el objetivo"),
    ("1", "valor", "propia", "Vⱼ / max V", "importancia del objetivo"),
    ("2", "p_vs_mejor", "propia", "pᵢⱼ − maxⱼ′ pᵢⱼ′ (≤ 0)", "¿es j el mejor destino posible para el arma i?"),
    ("3", "es_global", "flag", "1 solo en la fila 0", "marca la fila que construye la query de decisión"),
    ("4", "no_disponible", "flag", "1 si el arma i ya fue asignada, o si es la fila global", "MÁSCARA: logit −10⁹"),
    ("5", "asignado", "flag", "1 si xᵢⱼ = 1 en la solución parcial", "qué pares ya se eligieron (contexto)"),
    ("6", "q_sobrevive", "relativa", "qⱼ", "cuánto queda por destruir en j"),
    ("7", "valor_en_juego", "relativa", "Vⱼ qⱼ / max V   (fila global: Σ V q / Σ V)", "valor que aún se puede ganar en j"),
    ("8", "ganancia ★", "relativa", "gᵢⱼ / max g sobre pares disponibles", "criterio de MMR (análogo a dist_a_actual)"),
    ("9", "es_mmr", "relativa", "1 si (i, j) es el par que elegiría MMR", "elección de la heurística clásica"),
    ("10", "armas_en_obj", "relativa", "nº de armas ya asignadas a j / W", "rendimientos decrecientes"),
    ("11", "ventaja", "relativa", "pᵢⱼ − media_{i′∈F} pᵢ′ⱼ", "¿es esta arma comparativamente buena para j?"),
    ("12", "arrepentimiento", "relativa", "(gᵢⱼ − max_{i′∈F, i′≠i} gᵢ′ⱼ) / max g", "¿otra arma libre aprovecharía mejor j?"),
    ("13", "mmr_completa ★", "relativa", "1 si al completar con MMR el arma i va a j", "anticipación de decisiones futuras"),
    ("14", "q_final_mmr", "relativa", "qⱼ al completar con MMR", "¿quedará j bien cubierto al final?"),
    ("15", "frac_restantes", "global", "|F| / W (repetida en todas las filas)", "etapa de la construcción"),
]
rows = [("#", "Columna", "Rol", "Contenido", "Por qué ayuda a decidir")]
rows += [(n, [(nm, {"font": "Liberation Mono", "size": 8.5})], role, cont, why) for n, nm, role, cont, why in FEAT]
table(rows, [0.8, 3.1, 1.7, 5.9, 5.5], fills=[ROLE_FILL[f[2]] for f in FEAT], size=8.5,
      aligns=[C, WD_ALIGN_PARAGRAPH.LEFT, C, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
bullet([b("Query de decisión: "), ("la fila marcada con ", {}), c("es_global"),
        (" (columna 3). En el WTAP no hay una “ciudad actual”: el estado es la solución parcial completa. Tras la "
         "self-attention, la fila global concentra información de todos los pares. La query se completa con el promedio "
         "de los embeddings de los pares disponibles (sección 3).", {})])
bullet([b("Máscara: "), c("no_disponible"), (" (columna 4). Vale 1 en la fila global y en los T pares de cada arma ya "
        "asignada. Así, con k armas libres quedan exactamente k·T filas válidas, las mismas acciones que genera el entorno.",
                                             {})])
bullet([b("Información global: "), ("se repite en todas las filas (", {}), c("frac_restantes"),
        (") y también se guarda en la fila especial (", {}), c("valor_en_juego"), (" total). Las características relativas "
        "solo se calculan para pares disponibles; en las filas enmascaradas valen 0.", {})])
bullet([b("Independencia del índice y del tamaño: "), ("ninguna columna contiene el número de arma ni de objetivo, y todas "
        "están normalizadas por instancia (divididas por el máximo o expresadas como fracciones). Por eso el mismo modelo "
        "procesa 81, 151 o 301 filas. La relación entre pares (por ejemplo, otra arma mejor para el mismo objetivo) "
        "puede además aprenderse con la self-attention, porque el par (i, j) puede atender a los pares (i′, j).", {})])
figure(os.path.join(FIG, "estado.png"), TEXT_W,
       f"Figura 2. state2vec del paso {D['step']} del ejemplo (matriz transpuesta: cada columna del gráfico es una fila de "
       "la matriz; G = fila global; luego los 8 objetivos de cada arma). En rojo, la fila que escoge el experto. Cada "
       "característica se reescaló a [0, 1] solo para visualizarla.")

# ================================================================ 3. Salida
doc.add_heading("3. Codificación de la salida", level=1)
para([("El modelo entrega un ", {}), b(f"vector de logits de largo n = 1 + W·T"),
      (" (uno por fila) y un softmax lo convierte en la probabilidad de que cada fila sea la próxima acción:", {})])
formula(r"$u_r=\frac{(W_Q^{ptr}\,q)\cdot(W_K^{ptr}\,h_r)}{\sqrt{d}},\qquad u_r=-10^9\ \mathrm{si}\ x_{r,4}=1,"
        r"\qquad P(r)=\mathrm{softmax}(u)_r$", "puntero", height_cm=1.15)
bullet([b("Acción ↔ fila: "), ("la acción ", {}), c("(\"constructive\", (i, j))"), (" es “elegir la fila r = 1 + i·T + j”, y "
        "a la inversa i = (r − 1) // T, j = (r − 1) % T. Aunque la acción involucra dos elementos (un arma y un objetivo), "
        "cada par tiene su propia fila, así que el puntero lo selecciona directamente sin combinar filas, como sí haría "
        "falta en un movimiento 2-opt. El agente consulta ", {}), c("probs[1 + i·T + j]"), (" para cada acción legal.", {})])
bullet([b("Filas enmascaradas: "), ("las que tienen ", {}), c("no_disponible = 1"),
        (" (columna 4), es decir, la fila global y todos los pares de armas ya asignadas. Reciben logit −10⁹ antes del "
         "softmax, por lo que su probabilidad es exactamente 0 y no compiten con las acciones válidas.", {})])
bullet([b("Etiqueta Y: "), ("el índice entero de la fila escogida por el experto, Y = 1 + i*·T + a*ᵢ*. Equivale a un "
        "vector one-hot de largo n (1 en la acción escogida y 0 en el resto), pero se guarda como entero (tensor ", {}),
        c("long"), (") para usar ", {}), c("CrossEntropyLoss"), (":", {})])
formula(r"$\mathcal{L}=-\frac{1}{N}\sum_{s=1}^{N}\log P_s(Y_s), \qquad Y_s = 1+i^\star T+a^*_{i^\star}$", "loss",
        height_cm=1.05)
para([("Por la máscara, el softmax se reparte solo entre las k·T acciones válidas; un modelo sin entrenar parte con una "
       "pérdida cercana a log(k·T) en promedio. Como la etiqueta nunca está enmascarada (se verifica con un assert sobre "
       "todas las muestras), la pérdida siempre es finita.", {})])
doc.add_heading("Cómo se construye q: adaptación de TSP_Attention", level=2)
para([("La estructura del notebook se mantiene (embedding compartido → self-attention con residual → feed-forward → "
       "query → puntero). Cambian la query y la máscara:", {})], keep=True)
formula(r"$q = W_c\,[\,h_{\mathrm{global}}\ ;\ \bar h_{\mathrm{disp}}\,], \qquad"
        r"\bar h_{\mathrm{disp}}=\frac{\sum_r (1-x_{r,4})\,h_r}{\sum_r (1-x_{r,4})}$", "query", height_cm=1.15)
rows = [("", "TSP_Attention (notebook)", "WTAP_Attention (esta tarea)"),
        ("Filas", "ciudades", "fila global + pares (arma, objetivo)"),
        ("VEC_LEN", "6", "16"),
        ("Encoder", "1 bloque", "2 bloques con LayerNorm y Dropout, d = 64"),
        ("Query (paso 3)", "Wc[h_actual ; h_inicial] (col. 2 y fila 0)", "Wc[h_global ; promedio de disponibles] (col. 3)"),
        ("Máscara (paso 4)", "visitada (col. 3)", "no_disponible (col. 4)"),
        ("Puntero", "(W_Q q)·(W_K h_i)/√d", "igual")]
table(rows, [3.2, 6.3, 7.4], size=9)
g_model = T_["Modelo (atención)"]["mean_gap"]
g_mmr = T_["MMR (heurística clásica)"]["mean_gap"]
para([b("Resultado (puntos opcionales 7–8, detallados en el notebook). "),
      (f"Entrenado con {miles(R['n_train_samples'])} pares (estado, fila) "
       f"de {miles(R['n_train_instances'])} instancias, el modelo coincide con el experto en " +
       f"{pct(100 * R['val_acc_final'])} de las decisiones de validación. Usado como política de un GreedyAgent, logra "
       f"un gap de {pct(g_model, 2)} respecto del experto en las mismas 50 instancias de prueba de la Tarea 4 "
       f"(MMR: {pct(g_mmr, 2)}; MLP de la Tarea 4: {pct(R['gap_mlp_t4'], 2)}), y funciona sin cambios en "
       "instancias más grandes (" + ", ".join(f"{k}: {pct(v['Modelo (atención)'])}" for k, v in R["generalization"].items())
       + ").", {})])
same_mmr = D["mmr"] == D["chosen"]
figure(os.path.join(FIG, "salida.png"), 7.8,
       "Figura 3. Salida del puntero entrenado en el estado de la Figura 2, reordenada como armas × objetivos. "
       "Gris = filas enmascaradas; rojo = acción del experto; naranja = MMR. "
       + (f"El modelo pone probabilidad {D['p_model']:.2f} ".replace(".", ",")
          + ("en el par del experto" if D["model"] == D["chosen"] else f"en arma {D['model'][0]} → objetivo {D['model'][1]}")
          + (", que aquí coincide con MMR." if same_mmr else ".")))

doc.add_heading("Referencias", level=2)
for ref in [
    "den Broeder, G. G., Ellison, R. E., & Emerling, L. (1959). On optimum target assignments. Operations Research, 7(3).",
    "Lloyd, S. P., & Witsenhausen, H. S. (1986). Weapons allocation is NP-complete. Summer Computer Simulation Conference.",
    "Ahuja, R. K., Kumar, A., Jha, K. C., & Orlin, J. B. (2007). Exact and heuristic algorithms for the weapon-target "
    "assignment problem. Operations Research, 55(6).",
    "Vinyals, O., Fortunato, M., & Jaitly, N. (2015). Pointer Networks. NeurIPS.",
    "Kool, W., van Hoof, H., & Welling, M. (2019). Attention, Learn to Solve Routing Problems! ICLR.",
    "Notebook del curso: “Atención para el TSP: de la MLP al puntero”.",
]:
    para(ref, size=9, space_after=2)

doc.save(OUT)
for f in os.listdir(TMP):
    os.remove(os.path.join(TMP, f))
os.rmdir(TMP)
print("Informe guardado en", OUT)
